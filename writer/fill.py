# -*- coding: utf-8 -*-
"""
第 5 步 · 写回表格（fill.py）

把 matcher 抽出的值，按「第 i 条记录 → 待填区第 i 行」填回模板，
生成一份"已填写"的新 xlsx —— 这就是比赛要交的东西。

用法（VSCode 打开 writer 文件夹后，直接在 F5 下拉里选）：
    python fill.py                          # 默认模板4，正常写入
    python fill.py 模板3-会议任务跟踪表.xlsx   # 换别的模板
    python fill.py --overwrite              # 待填区已有内容时，覆盖写入（旧值会记进报告）

三条铁律（写在代码里，也写在报告里）：
  1. 永不改原模板 —— 输出永远是一份新文件（writer/输出/xxx-已填写.xlsx）
  2. 先查后写 —— 发现待填区已有内容时默认停下、列出冲突，由人拍板
  3. 只填值、不碰格式 —— 字体、边框、列宽、合并单元格全部原样保留
     （openpyxl 只改 cell.value 时，样式天然不动；这是它最省心的地方）
"""
import json
import os
import re
import sys
import time
from pathlib import Path

from openpyxl import load_workbook

# 让中文在终端正常显示（失败了也不影响）
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

HERE = Path(__file__).resolve().parent            # writer/
BASE = HERE.parent                                # 研一项目/
STRUCTURE_JSON = BASE / "template_reader" / "模板表结构.json"
MATCHER_DIR = BASE / "matcher"
# 模板目录：默认 测试数据/模板表格；可用环境变量 A23_TEMPLATE_DIR 覆盖
# （泛化测试 / 比赛换数据集时用；老流程不设变量时不受影响）
TEMPLATE_DIR = Path(os.environ.get("A23_TEMPLATE_DIR",
                                   str(BASE / "测试数据" / "模板表格")))
OUT_DIR = HERE / "输出"

DEFAULT_TEMPLATE = "模板4-客诉处理台账.xlsx"


def say(text=""):
    print(text)


def load_sheet_info(template_name):
    """从模板表结构.json 里找到这个模板的结构信息（第③层的产物）。"""
    data = json.loads(STRUCTURE_JSON.read_text(encoding="utf-8"))
    for item in data:
        if item.get("file") == template_name:
            return item["sheets"][0]
    return None


def load_extract(col_title):
    """读取某一列的抽取结果；没有就返回 None。"""
    f = MATCHER_DIR / ("抽取结果-%s.json" % col_title)
    if not f.exists():
        return None
    return json.loads(f.read_text(encoding="utf-8"))


def coerce(value):
    """把抽出的文本值转成合适的类型：纯数字 → 数字；其余原样。
    （这样"处理时长（天）"填进去的是数字 2，而不是文本"2"）"""
    s = str(value).strip()
    if re.fullmatch(r"-?\d+", s):
        return int(s)
    if re.fullmatch(r"-?\d+\.\d+", s):
        return float(s)
    return s


def scan_conflicts(ws, start_row, end_row, columns):
    """检查待填区里所有已有内容的格子。返回 [(行, 列标题, 坐标, 旧值), ...]"""
    found = []
    for r in range(start_row, end_row + 1):
        for col in columns:
            cell = ws.cell(row=r, column=col["col"])
            v = cell.value
            if v is not None and str(v).strip() != "":
                found.append((r, col["title"], cell.coordinate, v))
    return found


# ============================================================================
# Word（docx）模板的写回（2026-10-14 新增）
# ============================================================================
# 赛题评价方法：「再分别上传 5 个模版表格文件（word 或 excel）」
# Word 模板 = docx 里的一张表格（对应结构信息里的「表格1」= doc.tables[0]）。
#
# 与 xlsx 版的对应关系：
#     ws.cell(row, col).value = v    →    往 cell 的段落里加一个 run 文字
# 三条铁律不变：不改原模板 / 先查后写 / 只填值不碰格式。
#
# 写值方式选"段落里加 run"而不是"cell.text = v"：
#     后者会重建段落、丢掉模板给待填格设的对齐等格式；
#     add_run 保留段落格式，新文字随段落格式走（跟手打进去一样）。

def set_cell_text(cell, text):
    """往 docx 表格单元格写值：清掉旧文字（覆盖模式的残留），保留段落格式。"""
    p = cell.paragraphs[0]
    for run in list(p.runs):
        run.text = ""          # 保留 run 本身（=保留格式），只清文字
    p.add_run(text)


def fill_docx(template_name, tpl_path, overwrite):
    """Word 版写回：与 xlsx 版同规则、同报告格式。"""
    from docx import Document

    info = load_sheet_info(template_name)
    if not info:
        say("模板表结构.json 里没有这个模板：%s" % template_name)
        say("先在 template_reader 里把读取器跑一遍（python main.py \"模板目录\"），生成结构信息。")
        return 1

    columns = info["columns"]
    fr = info["fill_region"]
    start_row, end_row = fr["start_row"], fr["end_row"]

    say("=" * 62)
    say("写回表格（Word 版） · %s" % template_name)
    say("=" * 62)
    say("表格：%s ｜ 表头在第 %d 行 ｜ 共 %d 列" % (info["name"], info["header_row"], len(columns)))
    say("待填区：第 %d 行 ~ 第 %d 行（%d 行）" % (start_row, end_row, fr["row_count"]))
    say("")

    doc = Document(str(tpl_path))
    if not doc.tables:
        say("这个 docx 里没有表格，无法填写。")
        return 1
    table = doc.tables[0]          # 「表格1」→ 第 1 张表（结构信息里按序编号）

    # ---------- ① 先查 ----------
    conflicts = []
    for r in range(start_row, end_row + 1):
        for col in columns:
            v = (table.cell(r - 1, col["col"] - 1).text or "").strip()
            if v:
                conflicts.append((r, col["title"], "第%d行第%d列" % (r, col["col"]), v))

    say("① 检查待填区……")
    if conflicts:
        say("   发现 %d 个格子已有内容：" % len(conflicts))
        for r, title, coord, old in conflicts[:12]:
            say("     %s [%s] 旧值 = %s" % (coord, title, old))
        if len(conflicts) > 12:
            say("     …… 共 %d 个" % len(conflicts))
        if not overwrite:
            say("")
            say("已停下，没有写入任何东西。两条路：")
            say("  A. 打开模板，把上面这些内容清掉，然后重跑本脚本")
            say("  B. 确认可以覆盖的话，用 --overwrite 再跑一次（旧值会记进报告）")
            say("")
            say("（默认不覆盖，是怕把好数据冲掉；覆盖与否，由人拍板）")
            return 2
        say("   （--overwrite 模式：将覆盖以上内容，旧值会记进报告）")
    else:
        say("   干净：待填区没有任何已有内容。")
    say("")

    # ---------- ② 再写 ----------
    say("② 写入（第 i 条记录 → 待填区第 i 行）……")
    written = []          # (行, 列标题, 坐标, 值, 出处)
    missing_cols = []
    per_col_counts = {}   # 列名 -> 抽到几条（用于「序号」自动编号）

    for col in columns:
        res = load_extract(col["title"])
        if res is None:
            missing_cols.append(col["title"])
            continue
        items = res.get("items", [])
        per_col_counts[col["title"]] = len(items)
        for i, it in enumerate(items):
            r = start_row + i
            if r > end_row:
                break      # 超出待填区放不下了（报告里提示）
            cell = table.cell(r - 1, col["col"] - 1)
            set_cell_text(cell, str(coerce(it.get("value"))))
            src = ("%s %s" % (it.get("file") or "?", it.get("loc_text") or "")).strip()
            written.append((r, col["title"], "第%d行第%d列" % (r, col["col"]), it.get("value"), src))

    # 「序号」列自动编号：与 xlsx 版同规则（结构列，不来自源文档）
    seq_cols = [c for c in columns
                if c["title"] == "序号" and c["title"] not in per_col_counts]
    if seq_cols and per_col_counts:
        n_rec = max(per_col_counts.values())
        for col in seq_cols:
            for i in range(n_rec):
                r = start_row + i
                if r > end_row:
                    break
                cell = table.cell(r - 1, col["col"] - 1)
                set_cell_text(cell, str(i + 1))
                written.append((r, col["title"], "第%d行第%d列" % (r, col["col"]),
                                i + 1, "（序号·自动编号）"))
        say("   「序号」列自动编号 1~%d（结构列，不来自源文档）" % n_rec)

    missing_cols = [m for m in missing_cols if m != "序号"]
    if missing_cols:
        say("   注意：这些列没有抽取结果 → %s（先跑 matcher/run_all.py）" % "、".join(missing_cols))

    # ---------- ③ 保存（保存到新文件，原模板不动） ----------
    OUT_DIR.mkdir(exist_ok=True)
    out_path = OUT_DIR / (Path(template_name).stem + "-已填写.docx")
    try:
        doc.save(str(out_path))
    except PermissionError:
        say("")
        say("保存失败：输出文件正被占用，暂时写不进去 ——")
        say("  %s" % out_path)
        say("多半是它正在 Word / WPS 里开着。关掉那个窗口，然后重跑本脚本。")
        return 1

    # ---------- ④ 报告（与 xlsx 版同格式，文件名加 -word 区分） ----------
    lines = []
    lines.append("填表报告 · %s（Word 模板）" % template_name)
    lines.append("生成时间：%s" % time.strftime("%Y-%m-%d %H:%M"))
    lines.append("模式：%s" % ("覆盖写入（--overwrite）" if (overwrite and conflicts) else "正常写入"))
    lines.append("输出文件：%s" % out_path)
    lines.append("内容来源：matcher/抽取结果-*.json（每一格都能回溯到源文档）")
    lines.append("")
    lines.append("【一、冲突检查】")
    if conflicts:
        lines.append("待填区原有 %d 个格子有内容（%s）："
                     % (len(conflicts), "已被本次覆盖" if overwrite else "未处理"))
        for r, title, coord, old in conflicts:
            lines.append("  %s [%s] 旧值 = %s" % (coord, title, old))
    else:
        lines.append("干净：待填区没有任何已有内容。")
    lines.append("")
    lines.append("【二、写入明细】共 %d 格" % len(written))
    by_row = {}
    for r, title, coord, val, src in written:
        by_row.setdefault(r, []).append((title, coord, val, src))
    for r in sorted(by_row):
        lines.append("第 %d 行（第 %d 条记录）" % (r, r - start_row + 1))
        for title, coord, val, src in by_row[r]:
            lines.append("  %s = %s   [%s]   ← 出处：%s" % (title, val, coord, src))
        lines.append("")
    lines.append("【三、统计】")
    n_cols_done = len(columns) - len(missing_cols)
    lines.append("写入 %d 格（%d 列 × 最多 %d 条记录）" % (len(written), n_cols_done, len(by_row)))
    if conflicts and overwrite:
        lines.append("覆盖 %d 格（模板里原有的内容）" % len(conflicts))
    used_end = (max(by_row) if by_row else start_row - 1)
    if used_end < end_row:
        lines.append("第 %d ~ %d 行留空（源文档只有 %d 条记录，正常）"
                     % (used_end + 1, end_row, len(by_row)))
    lines.append("")
    lines.append("下一步：跑 verify_fill.py 验收（值对不对 + 格式保没保）。")

    report_path = OUT_DIR / ("填表报告-%s-word.txt" % Path(template_name).stem)
    report_path.write_text("\n".join(lines), encoding="utf-8")

    # 控制台同步打一份
    say("")
    say("\n".join(lines))
    say("")
    say("已生成：%s" % out_path)
    say("填表报告：%s" % report_path)
    return 0


def main():
    args = sys.argv[1:]
    overwrite = "--overwrite" in args
    plain = [a for a in args if not a.startswith("--")]
    template_name = plain[0] if plain else DEFAULT_TEMPLATE

    tpl_path = TEMPLATE_DIR / template_name
    if not tpl_path.exists():
        say("找不到模板文件：%s" % tpl_path)
        return 1

    # Word（docx）模板 → 走 Word 版写回路径（规则相同，写入方式不同）
    if tpl_path.suffix.lower() == ".docx":
        return fill_docx(template_name, tpl_path, overwrite)

    info = load_sheet_info(template_name)
    if not info:
        say("模板表结构.json 里没有这个模板：%s" % template_name)
        say("先在 template_reader 里把读取器跑一遍，生成结构信息。")
        return 1

    sheet_name = info["name"]
    columns = info["columns"]
    fr = info["fill_region"]
    start_row, end_row = fr["start_row"], fr["end_row"]

    say("=" * 62)
    say("写回表格 · %s" % template_name)
    say("=" * 62)
    say("工作表：%s ｜ 表头在第 %d 行 ｜ 共 %d 列" % (sheet_name, info["header_row"], len(columns)))
    say("待填区：第 %d 行 ~ 第 %d 行（%d 行）" % (start_row, end_row, fr["row_count"]))
    say("")

    # ---------- ① 先查 ----------
    wb = load_workbook(tpl_path)
    if sheet_name not in wb.sheetnames:
        say("模板里没有工作表「%s」，实际有：%s" % (sheet_name, wb.sheetnames))
        return 1
    ws = wb[sheet_name]

    conflicts = scan_conflicts(ws, start_row, end_row, columns)

    say("① 检查待填区……")
    if conflicts:
        say("   发现 %d 个格子已有内容：" % len(conflicts))
        for r, title, coord, old in conflicts[:12]:
            say("     %s [%s] 旧值 = %s" % (coord, title, old))
        if len(conflicts) > 12:
            say("     …… 共 %d 个" % len(conflicts))
        if not overwrite:
            say("")
            say("已停下，没有写入任何东西。两条路：")
            say("  A. 打开模板，把上面这些内容清掉，然后重跑本脚本")
            say("  B. 确认可以覆盖的话，用 --overwrite 再跑一次（旧值会记进报告）")
            say("")
            say("（默认不覆盖，是怕把好数据冲掉；覆盖与否，由人拍板）")
            return 2
        say("   （--overwrite 模式：将覆盖以上内容，旧值会记进报告）")
    else:
        say("   干净：待填区没有任何已有内容。")
    say("")

    # ---------- ② 再写 ----------
    say("② 写入（第 i 条记录 → 待填区第 i 行）……")
    written = []          # (行, 列标题, 坐标, 值, 出处)
    missing_cols = []
    per_col_counts = {}   # 列名 -> 抽到几条（用于「序号」自动编号）

    for col in columns:
        res = load_extract(col["title"])
        if res is None:
            missing_cols.append(col["title"])
            continue
        items = res.get("items", [])
        per_col_counts[col["title"]] = len(items)
        for i, it in enumerate(items):
            r = start_row + i
            if r > end_row:
                break      # 超出待填区放不下了（报告里提示）
            cell = ws.cell(row=r, column=col["col"])
            cell.value = coerce(it.get("value"))
            src = ("%s %s" % (it.get("file") or "?", it.get("loc_text") or "")).strip()
            written.append((r, col["title"], cell.coordinate, it.get("value"), src))

    # 「序号」列自动编号：序号是"表的结构"，不来自源文档 ——
    # 该列没有抽取结果时，按其它列的最大记录数自动填 1..N（比赛表格常见）
    seq_cols = [c for c in columns
                if c["title"] == "序号" and c["title"] not in per_col_counts]
    if seq_cols and per_col_counts:
        n_rec = max(per_col_counts.values())
        for col in seq_cols:
            for i in range(n_rec):
                r = start_row + i
                if r > end_row:
                    break
                cell = ws.cell(row=r, column=col["col"])
                cell.value = i + 1
                written.append((r, col["title"], cell.coordinate, i + 1,
                                "（序号·自动编号）"))
        say("   「序号」列自动编号 1~%d（结构列，不来自源文档）" % n_rec)

    missing_cols = [m for m in missing_cols if m != "序号"]   # 序号走自动编号，不算缺
    if missing_cols:
        say("   注意：这些列没有抽取结果 → %s（先跑 matcher/run_all.py）" % "、".join(missing_cols))

    # ---------- ③ 保存（保存到新文件，原模板不动） ----------
    OUT_DIR.mkdir(exist_ok=True)
    out_path = OUT_DIR / (Path(template_name).stem + "-已填写.xlsx")
    try:
        wb.save(out_path)
    except PermissionError:
        say("")
        say("保存失败：输出文件正被占用，暂时写不进去 ——")
        say("  %s" % out_path)
        say("多半是它正在 Excel / WPS 里开着。关掉那个窗口，然后重跑本脚本。")
        return 1

    # ---------- ④ 报告 ----------
    lines = []
    lines.append("填表报告 · %s" % template_name)
    lines.append("生成时间：%s" % time.strftime("%Y-%m-%d %H:%M"))
    lines.append("模式：%s" % ("覆盖写入（--overwrite）" if (overwrite and conflicts) else "正常写入"))
    lines.append("输出文件：%s" % out_path)
    lines.append("内容来源：matcher/抽取结果-*.json（每一格都能回溯到源文档）")
    lines.append("")
    lines.append("【一、冲突检查】")
    if conflicts:
        lines.append("待填区原有 %d 个格子有内容（%s）："
                     % (len(conflicts), "已被本次覆盖" if overwrite else "未处理"))
        for r, title, coord, old in conflicts:
            lines.append("  %s [%s] 旧值 = %s" % (coord, title, old))
    else:
        lines.append("干净：待填区没有任何已有内容。")
    lines.append("")
    lines.append("【二、写入明细】共 %d 格" % len(written))
    by_row = {}
    for r, title, coord, val, src in written:
        by_row.setdefault(r, []).append((title, coord, val, src))
    for r in sorted(by_row):
        lines.append("第 %d 行（第 %d 条记录）" % (r, r - start_row + 1))
        for title, coord, val, src in by_row[r]:
            lines.append("  %s = %s   [%s]   ← 出处：%s" % (title, val, coord, src))
        lines.append("")
    lines.append("【三、统计】")
    n_cols_done = len(columns) - len(missing_cols)
    lines.append("写入 %d 格（%d 列 × 最多 %d 条记录）" % (len(written), n_cols_done, len(by_row)))
    if conflicts and overwrite:
        lines.append("覆盖 %d 格（模板里原有的内容）" % len(conflicts))
    used_end = (max(by_row) if by_row else start_row - 1)
    if used_end < end_row:
        lines.append("第 %d ~ %d 行留空（源文档只有 %d 条记录，正常）"
                     % (used_end + 1, end_row, len(by_row)))
    lines.append("")
    lines.append("下一步：跑 verify_fill.py 验收（值对不对 + 格式保没保）。")

    report_path = OUT_DIR / ("填表报告-%s.txt" % Path(template_name).stem)
    report_path.write_text("\n".join(lines), encoding="utf-8")

    # 控制台同步打一份
    say("")
    say("\n".join(lines))
    say("")
    say("已生成：%s" % out_path)
    say("填表报告：%s" % report_path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
