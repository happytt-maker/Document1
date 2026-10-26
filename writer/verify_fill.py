# -*- coding: utf-8 -*-
"""
第 5 步 · 写回验收（verify_fill.py）

对照检查"已填写"文件，回答三个问题：

  1. 值对不对   —— ① 待填区与抽取结果逐格核对
                    ② 待填区以外必须纹丝未动（表头、说明行、标题……全表对比）
  2. 格式保没保 —— 单元格样式（字体/边框/对齐/填充/数字格式）、合并单元格、
                    列宽、行高、数据验证、条件格式，逐项对比
  3. 能不能交卷 —— 汇总结论

用法（VSCode 打开 writer 文件夹后，直接 F5）：
    python verify_fill.py                     # 默认模板4
    python verify_fill.py 模板3-会议任务跟踪表.xlsx

输出：writer/输出/验收报告-<模板名>.txt（控制台同步打印一份）
"""
import json
import os
import sys
import time
from pathlib import Path

from openpyxl import load_workbook

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

HERE = Path(__file__).resolve().parent            # writer/
BASE = HERE.parent                                # 研一项目/
STRUCTURE_JSON = BASE / "template_reader" / "模板表结构.json"
MATCHER_DIR = BASE / "matcher"
# 模板目录：默认 测试数据/模板表格；可用环境变量 A23_TEMPLATE_DIR 覆盖
# （泛化测试 / 比赛换数据集时用；老流程不设变量时不受影响）
TEMPLATE_DIR = Path(os.environ.get("A23_TEMPLATE_DIR",
                                   str(BASE / "测试数据" / "模板表格")))
OUT_DIR = HERE / "输出"

DEFAULT_TEMPLATE = "模板4-客诉处理台账.xlsx"


def say(text=""):
    print(text)


def load_sheet_info(template_name):
    data = json.loads(STRUCTURE_JSON.read_text(encoding="utf-8"))
    for item in data:
        if item.get("file") == template_name:
            return item["sheets"][0]
    return None


def load_extract(col_title):
    f = MATCHER_DIR / ("抽取结果-%s.json" % col_title)
    if not f.exists():
        return None
    return json.loads(f.read_text(encoding="utf-8"))


def norm(v):
    """值的规范化：None 和空串都算「空」，其余转字符串。"""
    if v is None:
        return ""
    return str(v).strip()


def style_key(cell):
    """把一个格子的"视觉样式"取成一串可比对的标签。"""
    f = cell.font
    b = cell.border
    a = cell.alignment
    fl = cell.fill
    return (
        f.name, f.size, bool(f.bold), str(f.color),
        b.left.style, b.right.style, b.top.style, b.bottom.style,
        a.horizontal, a.vertical, bool(a.wrap_text),
        str(cell.number_format),
        str(fl.patternType), str(fl.fgColor),
    )


def width_map(ws):
    m = {}
    for k, d in ws.column_dimensions.items():
        m[k] = d.width
    return m


def height_map(ws):
    m = {}
    for k, d in ws.row_dimensions.items():
        m[k] = d.height
    return m


# ============================================================================
# Word（docx）模板的验收（2026-10-14 新增）
# ============================================================================
# 与 xlsx 版同样回答三个问题：
#   1. 值对不对    —— 待填区 vs 抽取结果逐格核对 + 待填区以外纹丝未动（文字层面）
#   2. 格式保没保  —— Word 版的检查项：表格行列数 / 表格样式名 /
#                     待填格段落对齐（写入前后） / 正文段落数
#   3. 能不能交卷

def verify_docx(template_name):
    """Word 版写回验收：值对不对 + 待填区外有没有动 + 段落格式保没保。"""
    from docx import Document

    info = load_sheet_info(template_name)
    if not info:
        say("模板表结构.json 里没有这个模板：%s" % template_name)
        say("先在 template_reader 里把读取器跑一遍（python main.py \"模板目录\"），生成结构信息。")
        return 1

    columns = info["columns"]
    fr = info["fill_region"]
    start_row, end_row = fr["start_row"], fr["end_row"]

    orig_path = TEMPLATE_DIR / template_name
    filled_path = OUT_DIR / (Path(template_name).stem + "-已填写.docx")
    if not filled_path.exists():
        say("还没有「已填写」文件，先跑 fill.py：%s" % filled_path)
        return 1

    doc_o = Document(str(orig_path))
    doc_f = Document(str(filled_path))
    if not doc_o.tables or not doc_f.tables:
        say("模板或产物里没有表格，无法验收。")
        return 1
    table_o = doc_o.tables[0]
    table_f = doc_f.tables[0]

    lines = []
    problems = []          # (类别, 描述)

    lines.append("写回验收报告 · %s（Word 模板）" % template_name)
    lines.append("生成时间：%s" % time.strftime("%Y-%m-%d %H:%M"))
    lines.append("原模板：%s" % orig_path)
    lines.append("已填写：%s" % filled_path)
    lines.append("")

    # ---------- 1. 值：待填区 vs 抽取结果（与 xlsx 版同规则） ----------
    expected = {}          # (行, 列) -> 期望值
    n_expect = 0
    per_col_counts = {}    # 列名 -> 抽到几条（用于「序号」自动编号，与 fill.py 同规则）
    for col in columns:
        res = load_extract(col["title"])
        items = (res or {}).get("items", [])
        if items:
            per_col_counts[col["title"]] = len(items)
        for i, it in enumerate(items):
            r = start_row + i
            if r <= end_row:
                expected[(r, col["col"])] = norm(it.get("value"))
                n_expect += 1

    if per_col_counts:
        n_rec = max(per_col_counts.values())
        for col in columns:
            if col["title"] == "序号" and col["title"] not in per_col_counts:
                for i in range(n_rec):
                    r = start_row + i
                    if r <= end_row:
                        expected[(r, col["col"])] = str(i + 1)
                        n_expect += 1

    value_ok = 0
    value_bad = []
    n_fill_cells = 0
    for r in range(start_row, end_row + 1):
        for col in columns:
            n_fill_cells += 1
            exp = expected.get((r, col["col"]), "")
            got = norm(table_f.cell(r - 1, col["col"] - 1).text)
            if exp == got:
                value_ok += 1
            else:
                value_bad.append((r, col["title"], exp, got))

    lines.append("【1. 值 · 待填区】")
    lines.append("  与抽取结果一致：%d / %d 格（含未用到、应为空的格子）" % (value_ok, n_fill_cells))
    if value_bad:
        lines.append("  以下格子与期望不一致：")
        for r, title, exp, got in value_bad:
            lines.append("    第 %d 行 [%s] 期望 = %s ｜ 实际 = %s"
                         % (r, title, exp if exp else "（空）", got if got else "（空）"))
        problems.append(("值", "待填区有 %d 格不一致" % len(value_bad)))
    lines.append("")

    # ---------- 2. 值：待填区以外，必须纹丝未动 ----------
    n_rows = max(len(table_o.rows), len(table_f.rows))
    n_cols = max(len(table_o.columns), len(table_f.columns))
    diff_other = []
    for r in range(1, n_rows + 1):
        if start_row <= r <= end_row:
            continue
        for c in range(1, n_cols + 1):
            vo = norm(table_o.cell(r - 1, c - 1).text)
            vf = norm(table_f.cell(r - 1, c - 1).text)
            if vo != vf:
                diff_other.append(("第%d行第%d列" % (r, c), vo, vf))

    lines.append("【2. 值 · 待填区以外（表头/标题/说明行等）】")
    if diff_other:
        lines.append("  发现 %d 处被改动（不应该发生）：" % len(diff_other))
        for coord, vo, vf in diff_other:
            lines.append("    %s：原 = %s ｜ 现 = %s"
                         % (coord, vo if vo else "（空）", vf if vf else "（空）"))
        problems.append(("值", "待填区以外有 %d 处被改动" % len(diff_other)))
    else:
        lines.append("  全表其余格子一致，未被改动。")
    lines.append("")

    # ---------- 3. 格式（Word 版检查项） ----------
    fmt_bad = []

    size_o = (len(table_o.rows), len(table_o.columns))
    size_f = (len(table_f.rows), len(table_f.columns))
    if size_o != size_f:
        fmt_bad.append("表格行列数变了：原 %d×%d ｜ 现 %d×%d"
                       % (size_o[0], size_o[1], size_f[0], size_f[1]))

    style_o = getattr(table_o.style, "name", None)
    style_f = getattr(table_f.style, "name", None)
    if style_o != style_f:
        fmt_bad.append("表格样式变了：原 %s ｜ 现 %s" % (style_o, style_f))

    align_bad = []
    for r in range(start_row, end_row + 1):
        for col in columns:
            po = table_o.cell(r - 1, col["col"] - 1).paragraphs[0].alignment
            pf = table_f.cell(r - 1, col["col"] - 1).paragraphs[0].alignment
            if po != pf:
                align_bad.append("第%d行第%d列" % (r, col["col"]))
    if align_bad:
        fmt_bad.append("待填格段落对齐变了：%s%s"
                       % ("、".join(align_bad[:10]), "……" if len(align_bad) > 10 else ""))

    n_po = len(doc_o.paragraphs)
    n_pf = len(doc_f.paragraphs)
    if n_po != n_pf:
        fmt_bad.append("正文段落数变了：原 %d ｜ 现 %d" % (n_po, n_pf))

    lines.append("【3. 格式】")
    if fmt_bad:
        for b in fmt_bad:
            lines.append("  " + b)
        problems.append(("格式", "；".join(fmt_bad)))
    else:
        lines.append("  表格行列数：一致（%d×%d）" % size_o)
        lines.append("  表格样式：一致（%s）" % (style_o if style_o else "无样式"))
        lines.append("  待填格段落对齐：一致（写值未动对齐格式）")
        lines.append("  正文段落数：一致（%d 段）" % n_po)
    lines.append("")

    # ---------- 4. 结论 ----------
    lines.append("【4. 结论】")
    if not problems:
        lines.append("  全部通过 —— 值填对了，格式一点没动，可以交卷。")
    else:
        lines.append("  发现 %d 类问题，需要处理：" % len(problems))
        for cat, desc in problems:
            lines.append("    - [%s] %s" % (cat, desc))

    report_path = OUT_DIR / ("验收报告-%s-word.txt" % Path(template_name).stem)
    report_path.write_text("\n".join(lines), encoding="utf-8")

    say("")
    say("\n".join(lines))
    say("")
    say("验收报告：%s" % report_path)
    return 0 if not problems else 2


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    template_name = args[0] if args else DEFAULT_TEMPLATE

    # Word（docx）模板 → 走 Word 版验收路径
    if Path(template_name).suffix.lower() == ".docx":
        return verify_docx(template_name)

    info = load_sheet_info(template_name)
    if not info:
        say("模板表结构.json 里没有这个模板：%s" % template_name)
        return 1

    sheet_name = info["name"]
    columns = info["columns"]
    fr = info["fill_region"]
    start_row, end_row = fr["start_row"], fr["end_row"]

    orig_path = TEMPLATE_DIR / template_name
    filled_path = OUT_DIR / (Path(template_name).stem + "-已填写.xlsx")
    if not filled_path.exists():
        say("还没有「已填写」文件，先跑 fill.py：%s" % filled_path)
        return 1

    wb_o = load_workbook(orig_path)
    wb_f = load_workbook(filled_path)
    ws_o = wb_o[sheet_name]
    ws_f = wb_f[sheet_name]

    lines = []
    problems = []          # (类别, 描述)

    lines.append("写回验收报告 · %s" % template_name)
    lines.append("生成时间：%s" % time.strftime("%Y-%m-%d %H:%M"))
    lines.append("原模板：%s" % orig_path)
    lines.append("已填写：%s" % filled_path)
    lines.append("")

    # ---------- 1. 值：待填区 vs 抽取结果 ----------
    expected = {}          # (行, 列) -> 期望值
    n_expect = 0
    per_col_counts = {}    # 列名 -> 抽到几条（用于「序号」自动编号，与 fill.py 同规则）
    for col in columns:
        res = load_extract(col["title"])
        items = (res or {}).get("items", [])
        if items:
            per_col_counts[col["title"]] = len(items)
        for i, it in enumerate(items):
            r = start_row + i
            if r <= end_row:
                expected[(r, col["col"])] = norm(it.get("value"))
                n_expect += 1

    # 「序号」列：fill.py 会按记录数自动编号 1..N —— 验收期望值跟上同一规则
    if per_col_counts:
        n_rec = max(per_col_counts.values())
        for col in columns:
            if col["title"] == "序号" and col["title"] not in per_col_counts:
                for i in range(n_rec):
                    r = start_row + i
                    if r <= end_row:
                        expected[(r, col["col"])] = str(i + 1)
                        n_expect += 1

    value_ok = 0
    value_bad = []
    n_fill_cells = 0
    for r in range(start_row, end_row + 1):
        for col in columns:
            n_fill_cells += 1
            exp = expected.get((r, col["col"]), "")
            got = norm(ws_f.cell(row=r, column=col["col"]).value)
            if exp == got:
                value_ok += 1
            else:
                value_bad.append((r, col["title"], exp, got))

    lines.append("【1. 值 · 待填区】")
    lines.append("  与抽取结果一致：%d / %d 格（含未用到、应为空的格子）" % (value_ok, n_fill_cells))
    if value_bad:
        lines.append("  以下格子与期望不一致：")
        for r, title, exp, got in value_bad:
            lines.append("    第 %d 行 [%s] 期望 = %s ｜ 实际 = %s"
                         % (r, title, exp if exp else "（空）", got if got else "（空）"))
        problems.append(("值", "待填区有 %d 格不一致" % len(value_bad)))
    lines.append("")

    # ---------- 2. 值：待填区以外，必须纹丝未动 ----------
    diff_other = []
    max_r = max(ws_o.max_row, ws_f.max_row)
    max_c = max(ws_o.max_column, ws_f.max_column)
    for r in range(1, max_r + 1):
        if start_row <= r <= end_row:
            continue
        for c in range(1, max_c + 1):
            vo = norm(ws_o.cell(row=r, column=c).value)
            vf = norm(ws_f.cell(row=r, column=c).value)
            if vo != vf:
                coord = ws_o.cell(row=r, column=c).coordinate
                diff_other.append((coord, vo, vf))

    lines.append("【2. 值 · 待填区以外（表头/标题/说明行等）】")
    if diff_other:
        lines.append("  发现 %d 处被改动（不应该发生）：" % len(diff_other))
        for coord, vo, vf in diff_other:
            lines.append("    %s：原 = %s ｜ 现 = %s" % (coord, vo if vo else "（空）", vf if vf else "（空）"))
        problems.append(("值", "待填区以外有 %d 处被改动" % len(diff_other)))
    else:
        lines.append("  全表其余格子一致，未被改动。")
    lines.append("")

    # ---------- 3. 格式 ----------
    style_diffs = []
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            ko = style_key(ws_o.cell(row=r, column=c))
            kf = style_key(ws_f.cell(row=r, column=c))
            if ko != kf:
                style_diffs.append(ws_o.cell(row=r, column=c).coordinate)

    lines.append("【3. 格式】")
    if style_diffs:
        show = "、".join(style_diffs[:10]) + ("……" if len(style_diffs) > 10 else "")
        lines.append("  单元格样式差异：%d 处（%s）" % (len(style_diffs), show))
        problems.append(("格式", "样式差异 %d 处" % len(style_diffs)))
    else:
        lines.append("  单元格样式差异：0 处（字体/边框/对齐/填充/数字格式全一致）")

    # 合并单元格
    merged_o = {str(x) for x in ws_o.merged_cells.ranges}
    merged_f = {str(x) for x in ws_f.merged_cells.ranges}
    if merged_o == merged_f:
        lines.append("  合并单元格：一致（%s）" % ("、".join(sorted(merged_o)) if merged_o else "无"))
    else:
        lines.append("  合并单元格：不一致！原 = %s ｜ 现 = %s" % (sorted(merged_o), sorted(merged_f)))
        problems.append(("格式", "合并单元格不一致"))

    # 列宽 / 行高
    wo, wf = width_map(ws_o), width_map(ws_f)
    if wo == wf:
        lines.append("  列宽：一致")
    else:
        lines.append("  列宽：不一致！原 = %s ｜ 现 = %s" % (wo, wf))
        problems.append(("格式", "列宽不一致"))
    ho, hf = height_map(ws_o), height_map(ws_f)
    if ho == hf:
        lines.append("  行高：一致")
    else:
        lines.append("  行高：不一致！原 = %s ｜ 现 = %s" % (ho, hf))
        problems.append(("格式", "行高不一致"))

    # 数据验证 / 条件格式
    try:
        dv_o = len(ws_o.data_validations.dataValidation)
        dv_f = len(ws_f.data_validations.dataValidation)
        if dv_o == dv_f:
            lines.append("  数据验证（下拉等）：一致（%d 处）" % dv_o)
        else:
            lines.append("  数据验证：不一致！原 %d 处 / 现 %d 处" % (dv_o, dv_f))
            problems.append(("格式", "数据验证数量不一致"))
    except Exception as e:
        lines.append("  数据验证：读取异常（跳过）%s" % e)
    try:
        cf_o = sum(len(rules) for rules in ws_o.conditional_formatting._cf_rules.values())
        cf_f = sum(len(rules) for rules in ws_f.conditional_formatting._cf_rules.values())
        if cf_o == cf_f:
            lines.append("  条件格式：一致（%d 处）" % cf_o)
        else:
            lines.append("  条件格式：不一致！原 %d 处 / 现 %d 处" % (cf_o, cf_f))
            problems.append(("格式", "条件格式数量不一致"))
    except Exception as e:
        lines.append("  条件格式：读取异常（跳过）%s" % e)
    lines.append("")

    # ---------- 4. 结论 ----------
    lines.append("【4. 结论】")
    if not problems:
        lines.append("  全部通过 —— 值填对了，格式一点没动，可以交卷。")
    else:
        lines.append("  发现 %d 类问题，需要处理：" % len(problems))
        for cat, desc in problems:
            lines.append("    - [%s] %s" % (cat, desc))

    report_path = OUT_DIR / ("验收报告-%s.txt" % Path(template_name).stem)
    report_path.write_text("\n".join(lines), encoding="utf-8")

    say("")
    say("\n".join(lines))
    say("")
    say("验收报告：%s" % report_path)
    return 0 if not problems else 2


if __name__ == "__main__":
    sys.exit(main())
