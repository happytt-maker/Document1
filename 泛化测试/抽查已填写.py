# -*- coding: utf-8 -*-
"""
抽查已填写.py —— 打开三个场景的「已填写」交卷文件，把填写区原样打印出来。
（机器对分之外，肉眼终检一下：交卷的表格里到底长什么样）
"""
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent

FILES = [
    ("结果-场景1", "模板6-费用报销台账.xlsx", "费用报销台账"),
    ("结果-场景2", "模板7-候选人初筛表.xlsx", "候选人初筛"),
    ("结果-场景3", "模板8-设备报修工单登记表.xlsx", "报修工单登记"),
]

for scene_dir, tpl, sheet in FILES:
    stem = Path(tpl).stem
    path = ROOT / scene_dir / (stem + "-已填写.xlsx")
    print("=" * 70)
    print("%s ｜ %s" % (scene_dir, path.name))
    print("=" * 70)
    if not path.exists():
        print("  [找不到文件]")
        continue
    wb = load_workbook(str(path))
    ws = wb[sheet]
    for r in range(4, 10):          # 填写区 4~9 行
        vals = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            vals.append("—" if v is None or str(v).strip() == "" else str(v))
        print("  行%d  " % r + " ｜ ".join(vals))
    wb.close()
    print("")

print("（— = 空）")
