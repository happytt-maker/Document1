# -*- coding: utf-8 -*-
"""
跑五表.py —— 五表连测：模拟比赛现场"一套文档 × 5 张模板，逐张跑"

比赛的评价方法（赛题原文）：
    16 份文档全部上传 → 分别上传 5 张模板 → 每上传一张，系统跑一次 →
    记录响应时间；填完后与标准表格对比算准确率；每表准确率 ≥80%、每文档 ≤90 秒。

本脚本就是这场考试的预演：
    前置：读源文档（18 份，模拟"先把文档全部上传"，只做一次）
    逐张：模板理解 → 抽取（并行）→ 写回 → 验收（每张计时、归档、互不污染）

安全措施（沿用泛化测试的做法）：
    - 预检：5 张模板的待填区必须干净，不干净就停
    - 备份：跑前把 matcher / writer / doc_reader / template_reader 的相关产物备份
    - 还原：跑完自动恢复（模板4 老流程的状态原样如初）
    - 产物：每张表归档到「结果-模板X/」，互不覆盖

跑法（终端）：
    python 跑五表.py          # 全部 5 张
    python 跑五表.py 4        # 只跑某一张（编号 1~5，试水用）
    python 跑五表.py --no-preflight   # 跳过预检（不建议）
"""
import json
import os
import re
import shutil
import subprocess
import sys
import time
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

ROOT = Path(__file__).resolve().parent           # 五表连测/
BASE = ROOT.parent                               # 研一项目/
PY = sys.executable
# 备份目录：运行时由 backup_all() 赋值为「.备份-连测前-月日-时分秒」。
# 每次跑都新建一个带时间戳的目录、不删除旧备份 —— 因为"批量删除旧备份目录"
# 会触发宿主机的安全策略（safe-delete），连测进程会被直接杀掉（2026-10-11 踩坑）。
BACKUP = None

TPL_DIR = BASE / "测试数据" / "模板表格"
SRC_DIR = BASE / "测试数据" / "源文档"
MATCHER = BASE / "matcher"
WRITER = BASE / "writer"
WOUT = WRITER / "输出"
DR = BASE / "doc_reader"
TR = BASE / "template_reader"

TEMPLATES = [
    "模板1-设备采购汇总表.xlsx",
    "模板2-员工信息登记表.xlsx",
    "模板3-会议任务跟踪表.xlsx",
    "模板4-客诉处理台账.xlsx",
    "模板5-原材料库存汇总表.xlsx",
]


def run(cmd, cwd, env_extra=None, timeout=300):
    """跑一个子脚本，返回 (秒数, 返回码, 输出文本)。"""
    t0 = time.time()
    env = os.environ.copy()
    env["PYTHONIOENCODING"] = "utf-8"
    if env_extra:
        env.update(env_extra)
    try:
        p = subprocess.run([str(x) for x in cmd], cwd=str(cwd),
                           capture_output=True, env=env, timeout=timeout)
        out = p.stdout.decode("utf-8", "replace")
        err = p.stderr.decode("utf-8", "replace")
        if err.strip():
            out += "\n--- stderr ---\n" + err
        return time.time() - t0, p.returncode, out
    except subprocess.TimeoutExpired:
        return time.time() - t0, 999, "超时（%d 秒）" % timeout


def tail(text, n=12):
    lines = [ln for ln in (text or "").splitlines() if ln.strip()]
    return "\n".join("      | " + ln for ln in lines[-n:])


# ==================================================================
# 预检：5 张模板的待填区必须干净
# ==================================================================
def precheck():
    from openpyxl import load_workbook

    tinfo = json.loads((TR / "模板表结构.json").read_text(encoding="utf-8"))
    print("【预检】5 张模板的待填区是否干净（防止把旧内容当新数据）……")
    problems = []

    for tpl in TEMPLATES:
        info = next((t for t in tinfo if t.get("file") == tpl), None)
        if not info or not info.get("sheets"):
            problems.append("%s：模板表结构.json 里没有它，先跑 template_reader" % tpl)
            continue
        sheet = info["sheets"][0]
        fr = sheet.get("fill_region")
        if not fr:
            problems.append("%s：没找到待填区" % tpl)
            continue

        wb = load_workbook(str(TPL_DIR / tpl))
        ws = wb[sheet["name"]]
        dirty = []
        for r in range(fr["start_row"], fr["end_row"] + 1):
            for col in sheet["columns"]:
                v = ws.cell(row=r, column=col["col"]).value
                if v is not None and str(v).strip():
                    dirty.append("%s=%s" % (ws.cell(row=r, column=col["col"]).coordinate, v))
        wb.close()

        if dirty:
            problems.append("%s：待填区有 %d 格旧内容（%s）"
                            % (tpl, len(dirty), "、".join(dirty[:6])))
        else:
            print("  ✓ %s：干净（待填区 %d~%d 行）"
                  % (tpl, fr["start_row"], fr["end_row"]))

    if problems:
        print("")
        print("发现 %d 个问题 —— 先处理干净再跑：" % len(problems))
        for p in problems:
            print("  - " + p)
    return problems


# ==================================================================
# 备份 / 还原（跑完把工作区恢复成老样子）
# ==================================================================
def backup_all():
    """把跑前产物备份到「.备份-连测前-月日-时分秒」。

    注意：不删除旧备份目录（批量删除会触发宿主机安全策略、连测进程被杀）。
    旧备份留档不管，跑完的报告里会提示 —— 需要腾地方时手动删即可。
    """
    global BACKUP
    BACKUP = ROOT / (".备份-连测前-" + time.strftime("%m%d-%H%M%S"))
    BACKUP.mkdir(parents=True, exist_ok=False)

    n = 0
    for src, dst in [(DR / "output.json", BACKUP / "output.json"),
                     (TR / "模板表结构.json", BACKUP / "模板表结构.json"),
                     (MATCHER / "运行摘要.txt", BACKUP / "运行摘要.txt")]:
        if src.exists():
            shutil.copy2(src, dst)
            n += 1
    for f in MATCHER.glob("抽取结果-*.json"):
        shutil.copy2(f, BACKUP / f.name)
        n += 1
    for f in MATCHER.glob("评测报告-*.txt"):
        shutil.copy2(f, BACKUP / f.name)
        n += 1
    (BACKUP / "writer输出").mkdir(exist_ok=True)
    for f in WOUT.glob("*"):
        if f.is_file():
            shutil.copy2(f, BACKUP / "writer输出" / f.name)
            n += 1
    return n


def restore_all():
    # 先清掉连测期间生成的新产物
    for f in MATCHER.glob("抽取结果-*.json"):
        f.unlink()
    for f in MATCHER.glob("评测报告-*.txt"):
        f.unlink()
    for f in WOUT.glob("*"):
        if f.is_file():
            f.unlink()

    # 再从备份恢复
    for name, dst_dir in [("output.json", DR),
                          ("模板表结构.json", TR),
                          ("运行摘要.txt", MATCHER)]:
        src = BACKUP / name
        if src.exists():
            shutil.copy2(src, dst_dir / name)
    for f in BACKUP.glob("抽取结果-*.json"):
        shutil.copy2(f, MATCHER / f.name)
    for f in BACKUP.glob("评测报告-*.txt"):
        shutil.copy2(f, MATCHER / f.name)
    for f in (BACKUP / "writer输出").glob("*"):
        shutil.copy2(f, WOUT / f.name)


# ==================================================================
# 单张模板：跑一轮（模板 → 抽取 → 写回 → 验收），计时 + 归档
# ==================================================================
def run_one(tpl, no):
    out_dir = ROOT / ("结果-模板%d" % no)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = Path(tpl).stem
    stats = {"tpl": tpl, "no": no, "stages": {}}

    # 清掉上一轮的抽取结果 —— 保证本轮归档/对分用的全是本轮产物
    for f in MATCHER.glob("抽取结果-*.json"):
        f.unlink()

    # ② 理解模板（全量重扫；单张的模板理解其实是亚秒级）
    sec, code, out = run([PY, "main.py", str(TPL_DIR)], cwd=TR)
    stats["stages"]["模板"] = round(sec, 2)
    if code != 0:
        stats["error"] = "② 理解模板失败：\n" + tail(out, 12)
        return stats

    tinfo = json.loads((TR / "模板表结构.json").read_text(encoding="utf-8"))
    info = next((t for t in tinfo if t.get("file") == tpl), None)
    cols = [c["title"] for c in info["sheets"][0]["columns"]] if info else []
    stats["cols"] = cols

    # ③ 抽取（并行调模型）
    env = {"A23_TEMPLATE_NAME": tpl}
    sec, code, out = run([PY, "run_all.py", "--no-eval"], cwd=MATCHER, env_extra=env)
    stats["stages"]["抽取"] = round(sec, 2)
    if code != 0:
        stats["error"] = "③ 抽取失败：\n" + tail(out, 15)
        return stats

    # 逐列耗时（从运行摘要.txt 抓）
    summ = MATCHER / "运行摘要.txt"
    if summ.exists():
        txt = summ.read_text(encoding="utf-8")
        stats["col_lines"] = [ln for ln in txt.splitlines()
                              if "｜" in ln and "秒" in ln]

    # ④ 写回
    env = {"A23_TEMPLATE_DIR": str(TPL_DIR), "A23_TEMPLATE_NAME": tpl}
    sec, code, out = run([PY, "fill.py", tpl], cwd=WRITER, env_extra=env)
    stats["stages"]["写回"] = round(sec, 2)
    stats["fill_code"] = code
    stats["fill_tail"] = tail(out, 8)
    if code not in (0, 2):
        stats["error"] = "④ 写回失败（返回码 %d）：\n%s" % (code, tail(out, 12))
        return stats

    # ⑤ 验收
    sec, code, out = run([PY, "verify_fill.py", tpl], cwd=WRITER, env_extra=env)
    stats["stages"]["验收"] = round(sec, 2)
    stats["verify_code"] = code
    stats["verify_tail"] = tail(out, 10)

    # 模板4：趁抽取结果还是本轮，跑对分（这是我们自己的尺子，不计入响应时间）
    if "模板4" in tpl:
        sec, code, out = run([PY, "evaluate.py"], cwd=MATCHER)
        stats["stages"]["对分"] = round(sec, 2)
        rep = MATCHER / "评测报告-模板4.txt"
        if rep.exists():
            txt = rep.read_text(encoding="utf-8")
            stats["score_lines"] = [ln.strip() for ln in txt.splitlines()
                                    if ("值 ·" in ln or "出处 ·" in ln or "已评测" in ln)]

    # 响应时间 = 模板 + 抽取 + 写回 + 验收（对分是开发工具，不算）
    stats["resp_sec"] = round(sum(stats["stages"].get(k, 0)
                                  for k in ("模板", "抽取", "写回", "验收")), 2)

    # 写入格数（从填表报告抓）
    rep = WOUT / ("填表报告-%s.txt" % stem)
    if rep.exists():
        txt = rep.read_text(encoding="utf-8")
        m = re.search(r"【二、写入明细】共 (\d+) 格", txt)
        stats["written"] = int(m.group(1)) if m else None

    # 归档
    picks = [
        (WOUT / (stem + "-已填写.xlsx"), stem + "-已填写.xlsx"),
        (WOUT / ("验收报告-%s.txt" % stem), "验收报告.txt"),
        (WOUT / ("填表报告-%s.txt" % stem), "填表报告.txt"),
        (MATCHER / "运行摘要.txt", "运行摘要-抽取.txt"),
        (MATCHER / "评测报告-模板4.txt", "评测报告-模板4.txt"),
    ]
    for src_f, dst in picks:
        if src_f.exists() and ("模板4" in tpl or "评测报告" not in dst):
            shutil.copy2(src_f, out_dir / dst)
    for col in cols:
        f = MATCHER / ("抽取结果-%s.json" % col)
        if f.exists():
            shutil.copy2(f, out_dir / f.name)

    return stats


# ==================================================================
# 主流程
# ==================================================================
def main():
    args = sys.argv[1:]
    only = None
    for a in args:
        if a.isdigit():
            only = int(a)
    jobs = [(i + 1, t) for i, t in enumerate(TEMPLATES)
            if only is None or (i + 1) == only]

    print("=" * 64)
    print("五表连测 · 跑 %d 张模板 ｜ 时间 %s"
          % (len(jobs), time.strftime("%Y-%m-%d %H:%M")))
    print("=" * 64)
    print("模拟比赛节奏：文档全部上传（前置一次）→ 每张模板跑一次系统")

    # ---- 预检 ----
    if "--no-preflight" not in args:
        print("")
        if precheck():
            print("")
            print("预检不通过 → 停下。清干净模板再跑。")
            return
    else:
        print("（--no-preflight：跳过预检）")

    # ---- 备份 ----
    n = backup_all()
    print("")
    print("【备份】已备份 %d 个文件（output.json / 模板表结构 / 抽取结果 / writer输出）→ %s"
          % (n, BACKUP.name))
    olds = [p.name for p in ROOT.glob(".备份-连测前-*") if p.is_dir() and p != BACKUP]
    if olds:
        print("  （提示：历史备份 %d 个：%s —— 如嫌占地方可手动删除）"
              % (len(olds), "、".join(sorted(olds)[:4])))

    all_stats = []
    try:
        # ---- 前置：读源文档（一次）----
        print("")
        sec, code, out = run([PY, "main.py", str(SRC_DIR)], cwd=DR)
        if code != 0:
            print("！前置读取失败：\n" + tail(out, 12))
            return
        data = json.loads((DR / "output.json").read_text(encoding="utf-8"))
        n_doc = len(data.get("documents", []))
        n_blk = sum(d.get("block_count", 0) for d in data.get("documents", []))
        print("【前置】读源文档：%d 份 / %d 块 ｜ %.2f 秒（一次性，不计入单表响应时间）"
              % (n_doc, n_blk, sec))
        pre_sec = sec

        # ---- 逐张跑 ----
        for no, tpl in jobs:
            print("")
            print("-" * 64)
            print("▶ 第 %d 张：%s" % (no, tpl))
            print("-" * 64)
            st = run_one(tpl, no)
            all_stats.append(st)
            if "error" in st:
                print("  ！失败：" + st["error"])
            else:
                sg = st["stages"]
                print("  ②模板 %.2fs → ③抽取 %.2fs → ④写回 %.2fs → ⑤验收 %.2fs"
                      % (sg.get("模板"), sg.get("抽取"), sg.get("写回"), sg.get("验收")))
                print("  单表响应时间：%.2f 秒（红线 90 秒）｜ 写入 %s 格 ｜ 验收返回码 %d"
                      % (st["resp_sec"], st.get("written", "?"), st["verify_code"]))
                if st.get("score_lines"):
                    for ln in st["score_lines"]:
                        print("    [模板4对分] " + ln)
    finally:
        # ---- 还原（无论中途是否出错都还原）----
        restore_all()
        print("")
        print("【还原】已恢复老产物（模板4 流程状态不受影响）")

    # ---- 总报告 ----
    lines = []
    lines.append("=" * 64)
    lines.append("五表连测报告（%s）" % time.strftime("%Y-%m-%d %H:%M"))
    lines.append("=" * 64)
    lines.append("")
    lines.append("前置（一次性）：读源文档 %d 份 / %d 块 ｜ %.2f 秒" % (n_doc, n_blk, pre_sec))
    lines.append("响应时间口径：单表 = 模板理解 + 抽取 + 写回 + 验收（赛题红线 90 秒/文档）")
    lines.append("")
    lines.append("【逐张成绩】")
    lines.append("-" * 64)
    for st in all_stats:
        if "error" in st:
            lines.append("%s ｜ 失败：%s" % (st["tpl"], st["error"].splitlines()[0]))
            continue
        sg = st["stages"]
        lines.append("%s（%d 列）" % (st["tpl"], len(st.get("cols", []))))
        lines.append("  响应 %.2f 秒 = 模板 %.2f + 抽取 %.2f + 写回 %.2f + 验收 %.2f ｜ 写入 %s 格 ｜ 验收 %s"
                     % (st["resp_sec"], sg.get("模板"), sg.get("抽取"), sg.get("写回"),
                        sg.get("验收"), st.get("written", "?"),
                        "通过" if st["verify_code"] == 0 else ("有问题（返回码 %d）" % st["verify_code"])))
        if st.get("fill_code") == 2:
            lines.append("  注意：写回时发现待填区有内容（已按 --overwrite 外的方式停下，未写）")
    lines.append("")
    lines.append("【模板4 对分（唯一有标准答案的表）】")
    for st in all_stats:
        if st.get("score_lines"):
            for ln in st["score_lines"]:
                lines.append("  " + ln)
    lines.append("")
    lines.append("【逐列耗时明细】")
    for st in all_stats:
        if st.get("col_lines"):
            lines.append("  —— %s ——" % st["tpl"])
            for ln in st["col_lines"]:
                lines.append("    " + ln)
    lines.append("")
    lines.append("【红线对照】")
    over = [st for st in all_stats if "error" not in st and st["resp_sec"] > 90]
    n_ok = len([st for st in all_stats if "error" not in st])
    if over:
        lines.append("  超红线：%s" % "、".join(st["tpl"] for st in over))
    else:
        lines.append("  %d 张全部 ≤90 秒 ✓" % n_ok)
    lines.append("=" * 64)
    report = "\n".join(lines)
    (ROOT / "五表连测报告.txt").write_text(report, encoding="utf-8")

    print("")
    print(report)
    print("（报告已存：五表连测报告.txt）")


if __name__ == "__main__":
    main()
