# -*- coding: utf-8 -*-
"""端到端：模块3 表格填写全链路 + FAQ 数量确认"""
import http.cookiejar
import json
import os
import time
import urllib.error
import urllib.parse
import urllib.request

BASE = "http://127.0.0.1:8000"
OUT = r"C:\Users\覃敬涛\WorkBuddy\2026-09-24-10-48-03\_e2e.txt"
lines = []
ok = 0
bad = 0


def chk(name, cond, detail=""):
    global ok, bad
    if cond:
        ok += 1
        lines.append("  [OK]   %-34s %s" % (name, detail))
    else:
        bad += 1
        lines.append("  [FAIL] %-34s %s" % (name, detail))


class C(object):
    def __init__(self):
        self.jar = http.cookiejar.CookieJar()
        self.op = urllib.request.build_opener(urllib.request.ProxyHandler({}),
                                              urllib.request.HTTPCookieProcessor(self.jar))

    def call(self, path, method="GET", data=None, timeout=300):
        body = None
        h = {}
        if data is not None:
            body = json.dumps(data, ensure_ascii=False).encode("utf-8")
            h["Content-Type"] = "application/json"
        req = urllib.request.Request(BASE + path, data=body, headers=h, method=method)
        try:
            r = self.op.open(req, timeout=timeout)
            return r.status, r.read()
        except urllib.error.HTTPError as e:
            return e.code, e.read()
        except Exception as e:
            return -1, ("%s: %s" % (type(e).__name__, e)).encode()

    def js(self, path, method="GET", data=None, timeout=300):
        code, b = self.call(path, method, data, timeout)
        try:
            return code, json.loads(b.decode("utf-8"))
        except Exception:
            return code, {}


c = C()
code, j = c.js("/api/auth/login", "POST",
               {"username": "demo_ai", "password": "123456", "portal": "user"})
chk("登录", code == 200, "HTTP %s" % code)

# FAQ 真实字段
code, j = c.js("/api/support/faq")
items = j.get("items", [])
chk("FAQ 快捷指令", code == 200 and len(items) >= 6, "HTTP %s, %d 条（字段 items）" % (code, len(items)))
if items:
    lines.append("       示例：%s" % " / ".join(
        (x.get("title") or x.get("q") or str(x))[:16] for x in items[:4]))

# 模板清单
code, j = c.js("/api/fill/templates")
tpls = j.get("templates", [])
chk("模板清单", code == 200 and len(tpls) >= 5, "HTTP %s, %d 张" % (code, len(tpls)))

# 跑模块3：模板1
target = None
for t in tpls:
    if t.get("kind") == "xlsx":
        target = t["name"]
        break
lines.append("")
lines.append("=== 模块3 端到端：%s ===" % target)
t0 = time.time()
code, j = c.js("/api/fill", "POST", {"template": target})
jid = j.get("job_id")
chk("提交填写任务", code == 200 and bool(jid), "HTTP %s, job=%s, 状态=%s" % (code, jid, j.get("status")))

final = {}
if jid:
    last_log = ""
    for i in range(180):          # 最多等 180s
        code, jj = c.js("/api/job/%s" % jid)
        st = jj.get("status") or ""
        log = jj.get("log") or jj.get("message") or ""
        if isinstance(log, list):
            log = " | ".join(str(x) for x in log[-3:])
        if log and str(log) != last_log:
            last_log = str(log)
            lines.append("       [%5.1fs] %s" % (time.time() - t0, str(log)[:150]))
        if st in ("done", "error", "failed", "cancelled"):
            final = jj
            break
        time.sleep(1.0)

elapsed = time.time() - t0
st = final.get("status") or ""
chk("任务跑到结束", st in ("done", "error", "failed"),
    "状态=%s, 总耗时=%.1fs" % (st, elapsed))
chk("任务成功（done）", st == "done", "状态=%s" % st)

result = final.get("result") or final.get("output") or {}
lines.append("       结果摘要：%s" % json.dumps(result, ensure_ascii=False)[:400])

# 产物文件
stem = os.path.splitext(target)[0]
outdir = r"C:\Users\覃敬涛\Desktop\研一项目\writer\输出"
cands = [os.path.join(outdir, stem + "-已填写.xlsx"),
         os.path.join(outdir, stem + "-已填写.docx")]
hit = next((p for p in cands if os.path.exists(p)), None)
chk("产物文件已生成", bool(hit), (os.path.basename(hit) + " %.0f KB" % (os.path.getsize(hit) / 1024))
    if hit else "没找到")

# 读产物看是否真填了值
if hit and hit.endswith(".xlsx"):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(hit, data_only=True)
        ws = wb.active
        filled = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12)):
            for cell in row:
                if cell.value not in (None, ""):
                    filled.append((cell.coordinate, str(cell.value)[:22]))
        chk("产物里有内容（非空表）", len(filled) > 5, "非空单元格 %d 个" % len(filled))
        lines.append("       前 8 个单元格：%s"
                     % " | ".join("%s=%s" % (a, b) for a, b in filled[:8]))
        wb.close()
    except Exception as e:
        chk("产物可被 openpyxl 打开", False, str(e)[:120])
        lines.append("       ^ 注意：这是 writer 兼容性问题")

# 下载接口（模板名是中文，必须 URL 编码；不编码 urllib 会直接抛异常）
code, b = c.call("/api/fill/download?template=" + urllib.parse.quote(target))
chk("下载接口可用", code == 200 and len(b) > 5000, "HTTP %s, %d 字节" % (code, len(b)))

lines.append("")
lines.append("总计：%d 项通过，%d 项失败" % (ok, bad))
open(OUT, "w", encoding="utf-8").write("\n".join(lines))
print("PASS=%d FAIL=%d" % (ok, bad))
