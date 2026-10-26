# -*- coding: utf-8 -*-
"""
_验收-格式扩展版.py —— 模块（1）· 第 7 步验收：复合指令 + 格式能力扩展

覆盖范围：
  复合（multi）    docx 右对齐+加粗 ｜ xlsx 表头加粗+居中 ｜ 复合失败回滚（不留半成品）
  新格式操作       斜体 / 下划线 / 删除线 / 高亮（docx · xlsx · md 三档都验）
  取消格式         "把加粗去掉"（docx 双轮 / md 双轮）
  负向             txt 加下划线 → 必须明确拒绝（纯文本没有该写法）

流程：①记录源文件指纹 → ②逐条真跑（run_edit 直调，等价一句话链路）
      → ③读回产物验证属性 → ④产物归档（供人工打开核对）
      → ⑤核对源文件指纹（原件零改动）→ ⑥报告存盘

跑法（项目根目录）：
  python agent/_验收-格式扩展版.py
"""
import hashlib
import re
import shutil
import sys
from pathlib import Path

BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE / "agent"))
from edit import run_edit                 # noqa: E402
import docx                               # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX  # noqa: E402
from openpyxl import load_workbook        # noqa: E402

DOCS = BASE / "测试数据" / "源文档"
OUT = BASE / "agent" / "输出"
ARCH = OUT / "验收-格式扩展"

LOG = []


def log(s=""):
    LOG.append(str(s))
    print(s)


RESULTS = []


def check(item, cond, detail):
    RESULTS.append((item, bool(cond), detail))
    log("    [%s] %s ｜ %s" % ("PASS" if cond else "FAIL", item, detail))


def md5(p: Path) -> str:
    return hashlib.md5(p.read_bytes()).hexdigest()


def run(text, workcopy=None, context=None):
    log("-" * 72)
    log("指令：%s" % text)
    meta = {}
    ok, report = run_edit(text, DOCS, context=context, meta=meta, workcopy=workcopy)
    for ln in report.splitlines():
        ls = ln.strip()
        if ls.startswith(("【解析】", "【定位】", "【执行】", "【输出】", "【失败】",
                          "【回滚】", "【多操作】", "──")):
            log("  " + ls)
    log("  => ok=%s" % ok)
    return ok, meta


def archive(meta, newname):
    src = Path(meta.get("out_path") or "")
    if not src.exists():
        return None
    ARCH.mkdir(parents=True, exist_ok=True)
    dst = ARCH / newname
    if dst.exists():
        dst.unlink()
    shutil.move(str(src), str(dst))
    log("  （产物归档 → %s）" % dst.name)
    return dst


def cleanup(meta):
    p = meta.get("out_path")
    if p and Path(p).exists():
        Path(p).unlink()


def nonempty_runs(p):
    return [r for r in p.runs if (r.text or "").strip()]


def header_row(ws):
    best, bn = 1, -1
    for r in range(1, min(6, ws.max_row) + 1):
        n = sum(1 for c in range(1, ws.max_column + 1)
                if ws.cell(row=r, column=c).value is not None)
        if n > bn:
            best, bn = r, n
    return best


def main():
    print("=" * 70)
    print("第 7 步验收 · 复合指令 + 格式能力扩展（斜体/下划线/删除线/高亮/取消）")
    print("=" * 70)

    # ---- ① 源文件指纹 ----
    src_files = [
        DOCS / "docx" / "设备采购合同.docx",
        DOCS / "xlsx" / "员工信息表.xlsx",
        DOCS / "md" / "项目周报-第12周.md",
        DOCS / "txt" / "设备验收记录.txt",
    ]
    before = {str(p): md5(p) for p in src_files}

    # ============ 1. 复合：右对齐 + 加粗（docx，用户报障原句） ============
    log("")
    log("【1】复合（docx）：右对齐 + 加粗")
    ok, m = run("把设备采购合同的标题右对齐，并且加粗")
    if ok:
        d = docx.Document(m["out_path"])
        runs = nonempty_runs(d.paragraphs[0])
        check("1.1 标题对齐=RIGHT",
              d.paragraphs[0].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.RIGHT,
              "align=%s" % d.paragraphs[0].paragraph_format.alignment)
        check("1.2 标题全部 run=加粗",
              bool(runs) and all(r.font.bold for r in runs),
              "bold=%s" % [r.font.bold for r in runs])
        archive(m, "设备采购合同_格式结果_右对齐加粗.docx")

    # ============ 2. 下划线（docx） ============
    log("")
    log("【2】下划线（docx）")
    ok, m = run("把设备采购合同的标题加上下划线")
    if ok:
        d = docx.Document(m["out_path"])
        runs = nonempty_runs(d.paragraphs[0])
        check("2.1 标题全部 run=下划线",
              bool(runs) and all(r.font.underline for r in runs),
              "underline=%s" % [r.font.underline for r in runs])
        archive(m, "设备采购合同_格式结果_标题下划线.docx")

    # ============ 3. 斜体（docx） ============
    log("")
    log("【3】斜体（docx）")
    ok, m = run("把设备采购合同的标题设成斜体")
    if ok:
        d = docx.Document(m["out_path"])
        runs = nonempty_runs(d.paragraphs[0])
        check("3.1 标题全部 run=斜体",
              bool(runs) and all(r.font.italic for r in runs),
              "italic=%s" % [r.font.italic for r in runs])
        archive(m, "设备采购合同_格式结果_标题斜体.docx")

    # ============ 4. 删除线（docx，金额语义目标） ============
    log("")
    log("【4】删除线（docx·金额）")
    ok, m = run("把设备采购合同里所有的金额加上删除线")
    if ok:
        d = docx.Document(m["out_path"])
        hit = [r.text for p in d.paragraphs for r in p.runs
               if r.font.strike and any(k in (r.text or "") for k in ("￥", "人民币"))]
        check("4.1 金额 run=删除线（≥2 处）", len(hit) >= 2,
              "命中 strike run: %s" % hit[:3])
        archive(m, "设备采购合同_格式结果_金额删除线.docx")

    # ============ 5. 高亮（docx） ============
    log("")
    log("【5】高亮（docx）")
    ok, m = run("把设备采购合同的标题加黄色高亮")
    if ok:
        d = docx.Document(m["out_path"])
        runs = nonempty_runs(d.paragraphs[0])
        check("5.1 标题 run=黄色高亮",
              bool(runs) and all(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in runs),
              "highlight=%s" % [r.font.highlight_color for r in runs])
        archive(m, "设备采购合同_格式结果_标题高亮.docx")

    # ============ 6. 取消加粗（docx，双轮） ============
    log("")
    log("【6】取消加粗（docx·双轮：先加粗 → 把它的加粗去掉）")
    ok1, m1 = run("把设备采购合同的标题加粗")
    if ok1:
        ok2, m2 = run("把它的加粗去掉",
                      workcopy={"doc": m1["doc"], "path": m1["out_path"]},
                      context={"current_doc": m1["doc"], "last_hint": "标题"})
        if ok2:
            d = docx.Document(m2["out_path"])
            runs = nonempty_runs(d.paragraphs[0])
            check("6.1 取消后标题 bold 非真",
                  not any(r.font.bold for r in runs),
                  "bold=%s" % [r.font.bold for r in runs])
            archive(m2, "设备采购合同_格式结果_取消加粗.docx")
        cleanup(m1)

    # ============ 7. 下划线（xlsx） ============
    log("")
    log("【7】下划线（xlsx）")
    ok, m = run("给员工信息表里的王振国加上下划线")
    if ok:
        ws = load_workbook(m["out_path"]).active
        cell = None
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "王振国" in c.value:
                    cell = c
                    break
            if cell:
                break
        check("7.1 王振国单元格=下划线",
              cell is not None and cell.font.underline in (True, "single"),
              "underline=%s" % (cell.font.underline if cell else None))
        archive(m, "员工信息表_格式结果_王振国下划线.xlsx")

    # ============ 8. 高亮底色（xlsx） ============
    log("")
    log("【8】高亮底色（xlsx）")
    ok, m = run("把员工信息表的表头标黄")
    if ok:
        ws = load_workbook(m["out_path"]).active
        hr = header_row(ws)
        fills = [ws.cell(row=hr, column=c).fill for c in range(1, ws.max_column + 1)]
        okfill = all((f.fill_type == "solid" and str(f.start_color.rgb).endswith("FFFF00"))
                     for f in fills)
        check("8.1 表头行=黄色填充", okfill,
              "fill_type=%s rgb=%s" % (fills[0].fill_type, fills[0].start_color.rgb))
        archive(m, "员工信息表_格式结果_表头标黄.xlsx")

    # ============ 9. 复合（xlsx：表头加粗 + 居中） ============
    log("")
    log("【9】复合（xlsx）：表头加粗 + 居中")
    ok, m = run("把员工信息表的表头加粗，并且居中")
    if ok:
        ws = load_workbook(m["out_path"]).active
        hr = header_row(ws)
        c1 = ws.cell(row=hr, column=1)
        check("9.1 表头格=加粗", c1.font.bold is True, "bold=%s" % c1.font.bold)
        check("9.2 表头格=居中", c1.alignment.horizontal == "center",
              "horizontal=%s" % c1.alignment.horizontal)
        archive(m, "员工信息表_格式结果_表头加粗居中.xlsx")

    # ============ 10. 斜体（md 标题） ============
    log("")
    log("【10】斜体（md）")
    ok, m = run("把项目周报的标题设成斜体")
    if ok:
        first = next(ln for ln in Path(m["out_path"]).read_text(encoding="utf-8").splitlines()
                     if ln.strip())
        check("10.1 首行=斜体标记",
              bool(re.search(r"(?<!\*)\*[^*]+\*(?!\*)", first)), repr(first[:60]))
        archive(m, "项目周报-第12周_格式结果_标题斜体.md")

    # ============ 11. 删除线（md 词级） ============
    log("")
    log("【11】删除线（md）")
    ok, m = run("把项目周报里所有的'问题'加上删除线")
    if ok:
        text = Path(m["out_path"]).read_text(encoding="utf-8")
        check("11.1 『问题』=删除线标记（≥2 处）",
              text.count("~~问题~~") >= 2, "count=%d" % text.count("~~问题~~"))
        archive(m, "项目周报-第12周_格式结果_问题删除线.md")

    # ============ 12. 取消加粗（md，双轮） ============
    log("")
    log("【12】取消加粗（md·双轮）")
    ok1, m1 = run("把项目周报的标题加粗")
    if ok1:
        ok2, m2 = run("把标题的加粗去掉",
                      workcopy={"doc": m1["doc"], "path": m1["out_path"]},
                      context={"current_doc": m1["doc"], "last_hint": "标题"})
        if ok2:
            first = next(ln for ln in Path(m2["out_path"]).read_text(encoding="utf-8").splitlines()
                         if ln.strip())
            check("12.1 取消后首行无 ** 标记", "**" not in first, repr(first[:60]))
            archive(m2, "项目周报-第12周_格式结果_取消加粗.md")
        cleanup(m1)

    # ============ 13. 负向：txt 下划线（必须明确拒绝） ============
    log("")
    log("【13】负向：txt 加下划线 → 明确拒绝")
    ok, m = run("把设备验收记录的标题加上下划线")
    check("13.1 txt 下划线=明确失败", not ok, "ok=%s" % ok)

    # ============ 14. 负向：复合失败回滚（不留半成品） ============
    log("")
    log("【14】负向：复合第 2 步找不到目标 → 回滚不留产物")
    leftover = OUT / "员工信息表_格式结果.xlsx"
    if leftover.exists():
        leftover.unlink()
    ok, m = run("把员工信息表的表头加粗，并且把'邮箱'列加粗")
    check("14.1 失败且无半成品产物",
          (not ok) and (not leftover.exists()),
          "ok=%s 无产物=%s" % (ok, not leftover.exists()))

    # ---- ⑤ 源文件指纹核对 ----
    log("")
    log("【原件核对】")
    for p in src_files:
        same = md5(p) == before[str(p)]
        check("原件未改动：%s" % p.name, same, "md5 一致" if same else "被改动！")

    # ---- ⑥ 汇总 ----
    log("")
    log("=" * 70)
    npass = sum(1 for _, c, _ in RESULTS if c)
    log("验收汇总：%d/%d 通过" % (npass, len(RESULTS)))
    for item, c, detail in RESULTS:
        if not c:
            log("  FAIL ｜ %s ｜ %s" % (item, detail))
    log("=" * 70)
    log("产物归档目录：%s" % ARCH)

    (OUT / "格式扩展-验收报告.txt").write_text("\n".join(LOG), encoding="utf-8")
    print("\n完整报告已存：agent/输出/格式扩展-验收报告.txt")
    return 0 if npass == len(RESULTS) else 1


if __name__ == "__main__":
    sys.exit(main())
