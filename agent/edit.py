# -*- coding: utf-8 -*-
"""
edit.py —— 模块（1）· 第 2~4 步：编辑 / 格式 / 提取 执行器（说一句话 → 文档真的动了）

链路：一句话 → ①解析（parse_intent）→ ②定位（在文档里找到位置）
      → ③执行（编辑：替换/删除/插入；格式：加粗/字号/字体/颜色/对齐/行距/缩进；
                提取：把字段信息整理出来、不改文档）
      → ④输出新文件（原件一个字不动）

支持的文档：.docx（python-docx，保格式）/ .xlsx（openpyxl，保格式）/ .md / .txt（纯文本）
支持的编辑操作（第 2 步）：
  replace 替换：把 X 换成 Y
  delete  删除：删掉 X
  insert  插入：在指定位置加上 X
支持的格式与排版操作（第 3 步，细节见 format_ops.py）：
  排版：set_align 对齐 / set_spacing 行距 / set_indent 首行缩进
  格式：set_bold 加粗 / set_color 颜色 / set_font 字体 / set_size 字号
支持的内容提取（第 4 步，细节见 extract_ops.py）：
  extract 提取：字段规范化（合称拆分/别名/语义翻译）→ 复用抽取线 → 配对整理

两个"人话细节"（都是上一轮预留的边界）：
  1. 标签 vs 值 —— 用户说"把签订日期改成 2026年4月1日"，解析出的 find 是
     "签订日期"（标签）。文档里是"签订日期：2026年3月12日"。
     若直接把"签订日期"四个字替换掉，会变成"2026年4月1日：2026年3月12日"（错！）。
     正确逻辑：find 后面紧跟冒号 → find 是标签 → 替换冒号后面的"值"。
  2. 分隔符处理 —— 从名单"周立、吴晓敏、李博文、陈雨薇"里删"李博文"，
     不能留下"、、"；所以删除时自动连一个相邻的顿号一起删。

跑法（在项目根目录）：
  python agent/edit.py "把项目推进会会议纪要参会人员里的'李博文'删掉"
  python agent/edit.py "把设备采购合同里的签订日期改成 2026年4月1日"
  python agent/edit.py "把设备采购合同的标题居中"
  python agent/edit.py "把设备采购合同里所有的金额加粗"
  python agent/edit.py "把员工信息表里所有员工的姓名和工号整理出来"

输出：
  屏幕：解析 / 定位 / 执行 / 输出 四段报告
  文件：agent/输出/<文档名>_编辑结果.<ext> 或 <文档名>_格式结果.<ext>
        + 编辑报告.txt / 格式报告.txt
"""
import argparse
import json
import re
import sys
import time
from pathlib import Path

import docx

# 让本脚本能直接 import 同目录的 parse_intent / format_ops
sys.path.insert(0, str(Path(__file__).resolve().parent))
from parse_intent import parse_one, scan_docs, as_dir_list   # noqa: E402
from format_ops import apply_format             # noqa: E402

# 输出兜底：遇到终端表示不了的字符就替换，不让程序崩
try:
    sys.stdout.reconfigure(errors="replace")
except Exception:
    pass

BASE = Path(__file__).resolve().parents[1]                 # = 研一项目/
DEFAULT_DOCS_DIR = BASE / "测试数据" / "源文档"              # 默认文档库（标准 18 份）
IMPORT_DIR = BASE / "导入文档"                              # 模块2：用户导入的文档库
DEFAULT_DOCS_DIRS = [DEFAULT_DOCS_DIR, IMPORT_DIR]         # 默认文档库（双目录：标准 + 导入）
OUT_DIR = BASE / "agent" / "输出"                           # 编辑结果的输出目录

SEP_CHARS = "、，；,;"          # 列表分隔符（删词时连删、插词时补位）
END_PUNCT = "。！？；.!?;"      # 句末标点（插词时插在它前面）

LAST_RUN_LABEL = "执行"         # 最近一次执行的类别（编辑/格式），用于报告文件名


# ==================== 一、找文档 ====================

def find_doc_file(doc_name: str, docs_dirs):
    """在文档库（一个目录或多个目录）里按名字找文件：先精确匹配，再宽松包含匹配。

    模块（2）起文档库是双目录（标准 + 导入），传列表即可；顺序即优先级
    （精确同名时，列表靠前的目录先命中——标准库在前）。
    找不到返回 None。
    """
    if not doc_name:
        return None
    files = []
    for d in as_dir_list(docs_dirs):
        files += [p for p in d.rglob("*") if p.is_file()]

    # 1) 精确匹配文件名
    for p in files:
        if p.name == doc_name:
            return p

    # 2) 宽松匹配：用去扩展名的核心词互相包含
    stem = Path(doc_name).stem
    if not stem:
        return None
    cands = [p for p in files if stem in p.stem or p.stem in stem]
    if not cands:
        return None
    if len(cands) == 1:
        return cands[0]

    # 3) 多个候选时，优先和后缀一致的
    ext = Path(doc_name).suffix.lower()
    same_ext = [p for p in cands if p.suffix.lower() == ext]
    return (same_ext or cands)[0]


# ==================== 二、docx 保格式替换（核心算法） ====================

def replace_in_runs(p, find: str, repl: str) -> int:
    """在段落 p 里把 find 全部替换成 repl，返回替换次数。

    docx 里一段文字可能被拆成多个 run（一串连续的"同格式小段"）。
    要保持格式，替换时不能粗暴地清空重写整段，而要：
      · find 起点所在的 run：保留前面的字 + 写入新字（新字继承这个 run 的格式）
      · 被 find 完全盖住的中间 run：清空
      · find 终点所在的 run：保留后面的字
    """
    if not find:
        return 0
    count = 0
    search_from = 0
    while True:
        runs = p.runs
        if not runs:
            return count
        texts = [r.text or "" for r in runs]
        full = "".join(texts)
        idx = full.find(find, search_from)
        if idx == -1:
            return count
        end = idx + len(find)

        # 每个 run 在整段文本里的起点位置
        starts = []
        pos = 0
        for t in texts:
            starts.append(pos)
            pos += len(t)

        # find 起点所在 run（si）、终点所在 run（ei）
        si = ei = None
        for i in range(len(starts) - 1, -1, -1):
            if starts[i] <= idx:
                si = i
                break
        for i in range(len(starts) - 1, -1, -1):
            if starts[i] < end:
                ei = i
                break
        if si is None or ei is None:
            return count

        si_off = idx - starts[si]          # find 在起点 run 内的偏移
        ei_off = end - starts[ei]          # find 终点在终点 run 内的偏移

        if si == ei:
            t = texts[si]
            runs[si].text = t[:si_off] + repl + t[ei_off:]
        else:
            t = texts[si]
            runs[si].text = t[:si_off] + repl
            for i in range(si + 1, ei):
                runs[i].text = ""
            t = texts[ei]
            runs[ei].text = t[ei_off:]

        count += 1
        search_from = idx + len(repl)      # 从替换产物之后继续找，防"新词含旧词"死循环
    return count


def delete_from_paragraph(p, find: str):
    """删掉段落里所有 find（自动连一个相邻分隔符）。返回 (删除次数, 处理方式列表)。"""
    count = 0
    notes = []
    for _ in range(50):                    # 防御上限：一段里最多处理 50 处
        text = p.text
        i = text.find(find)
        if i == -1:
            break
        after = text[i + len(find): i + len(find) + 1]
        before = text[i - 1:i] if i > 0 else ""
        if after in SEP_CHARS:
            target, how = find + after, "连后分隔符"
        elif before in SEP_CHARS:
            target, how = before + find, "连前分隔符"
        else:
            target, how = find, "无相邻分隔符"
        n = replace_in_runs(p, target, "")
        if n == 0:
            break
        count += n
        notes.append(how)
    return count, notes


def append_to_paragraph(p, text_to_add: str) -> bool:
    """在段落末尾加上内容（自动补顿号；若末尾是句末标点则插在标点前）。"""
    runs = list(p.runs)
    if not runs:
        return False
    last = None
    for r in reversed(runs):
        if (r.text or ""):
            last = r
            break
    if last is None:
        last = runs[-1]
    t = last.text or ""
    if t and t[-1] in END_PUNCT:
        last.text = t[:-1] + "、" + text_to_add + t[-1]
    else:
        last.text = t + "、" + text_to_add
    return True


# ==================== 三、执行文档编辑 ====================

TITLE_WORDS = ("标题", "题目", "标题行", "文档标题", "文章标题", "首行")


def _looks_like_title_op(find: str, hint: str) -> bool:
    """识别「把标题改成 X」类指令：
    find 为空（或本身就是"标题"这类方位词）且定位提示指向标题。

    为什么需要它：模型对同一句话的解析会飘——有时给 find=""，
    有时把"标题"两个字当搜索词（find="标题"）。这是执行器的兜底归一化。"""
    title_hint = ("标题" in hint) or ("题目" in hint) or ("首行" in hint)
    if not find:
        return title_hint
    return (find in TITLE_WORDS) and title_hint


def pick_title_paragraph(paras):
    """挑"标题段"：优先 Title 样式；否则第一个非空段。找不到返回 None。"""
    for p in paras:
        style_name = (p.style.name if p.style is not None else "") or ""
        if style_name == "Title" and p.text.strip():
            return p
    for p in paras:
        if p.text.strip():
            return p
    return None


def edit_docx_paragraphs(doc, intent: dict):
    """对打开的 docx 执行编辑。返回 (改动明细, 是否成功, 说明)。

    改动明细：[{"para": 段号, "before": 改前, "after": 改后, "how": 方式}, ...]
    """
    op = intent.get("op")
    params = intent.get("params") or {}
    find = (params.get("find") or "").strip()
    repl = params.get("replace_with") or ""
    add = params.get("text") or ""
    hint = ((intent.get("locate") or {}).get("hint") or "").strip()

    paras = list(doc.paragraphs)
    changes = []

    # ---- insert：找 hint 所在段落，在末尾加上内容 ----
    if op == "insert":
        if not add:
            return changes, False, "要加的内容为空"
        targets = [p for p in paras if hint and hint in p.text]
        if not targets:
            return changes, False, "没找到含「%s」的段落" % hint
        p = targets[0]
        before = p.text
        if not append_to_paragraph(p, add):
            return changes, False, "目标段落是空段落，没法追加"
        changes.append({"para": paras.index(p), "before": before, "after": p.text,
                        "how": "在段末追加「%s」（自动补顿号）" % add})
        return changes, True, "定位词「%s」锁定 1 段" % hint

    # ---- replace / delete：先找含 find 的段落 ----
    # 「把标题改成 X」：整段替换标题（保首个非空 run 的格式）；失败则落回普通路径
    if op == "replace" and _looks_like_title_op(find, hint) and repl:
        t = pick_title_paragraph(paras)
        if t is not None:
            before = t.text
            if replace_in_runs(t, before, repl):
                changes.append({"para": paras.index(t), "before": before, "after": t.text,
                                "how": "整段替换标题（保留格式）"})
                return changes, True, "标题整体替换"
    if not find:
        return changes, False, "没有要处理的目标词（find 为空）"
    broad = [p for p in paras if find in p.text]
    if not broad:
        return changes, False, "文档段落里找不到「%s」" % find

    # locate.hint 能同时命中时，把范围缩小到"hint + find"都出现的段落
    narrow = [p for p in broad if hint and hint in p.text]
    targets = narrow if narrow else broad
    scope = ("定位词「%s」精确命中 %d 段" % (hint, len(targets))) if narrow else \
            ("定位词没命中，按全文范围处理 %d 段" % len(targets))

    for p in targets:
        before = p.text
        if op == "replace":
            # 标签模式：find 后紧跟冒号 → 替换冒号后的"值"
            m = re.search(re.escape(find) + r"\s*[：:]", before)
            if m:
                vs = m.end()
                ve = len(before)
                for i in range(vs, len(before)):
                    if before[i] in "，。；,;\n":
                        ve = i
                        break
                val = before[vs:ve].strip()
                if val:
                    n = replace_in_runs(p, val, repl)
                    if n:
                        changes.append({"para": paras.index(p), "before": before,
                                        "after": p.text,
                                        "how": "标签模式：「%s」后的值「%s」→「%s」（%d 处）"
                                               % (find, val, repl, n)})
                    continue
                # 标签后没值：跳过并说明
                changes.append({"para": paras.index(p), "before": before, "after": before,
                                "how": "「%s」是标签但后面没有值，跳过" % find})
                continue
            # 普通模式：直接把 find 换成 repl
            n = replace_in_runs(p, find, repl)
            if n:
                changes.append({"para": paras.index(p), "before": before,
                                "after": p.text,
                                "how": "普通模式：「%s」→「%s」（%d 处）" % (find, repl, n)})
        else:  # delete
            n, notes = delete_from_paragraph(p, find)
            if n:
                changes.append({"para": paras.index(p), "before": before,
                                "after": p.text,
                                "how": "删除「%s」（%s，共 %d 处）"
                                       % (find, "；".join(sorted(set(notes))), n)})

    if not changes:
        return changes, False, "找到了含「%s」的段落，但替换/删除没有生效" % find
    return changes, True, scope


# ==================== 四、md / txt 文本编辑 ====================

def edit_plain_text(path: Path, intent: dict):
    """对 md / txt 执行编辑。返回 (改动明细, 是否成功, 说明, 新文本)。

    返回契约与 edit_docx_paragraphs 的前三位一致（改动/是否成功/说明），
    新文本殿后——md/txt 由调用方统一写回文件。"""
    op = intent.get("op")
    params = intent.get("params") or {}
    find = (params.get("find") or "").strip()
    repl = params.get("replace_with") or ""
    add = params.get("text") or ""
    hint = ((intent.get("locate") or {}).get("hint") or "").strip()

    text = path.read_text(encoding="utf-8")
    changes = []
    msg = "文本编辑"

    if op == "replace":
        done = False
        # 「把标题改成 X」：整行替换标题（保留 md 的 # 标记）；失败则落回普通路径
        if _looks_like_title_op(find, hint) and repl:
            lines = text.splitlines(keepends=True)
            hit = None
            for i, ln in enumerate(lines):
                if ln.strip():
                    hit = i
                    break
            if hit is not None:
                old_line = lines[hit].rstrip("\n")
                m = re.match(r"^(\s*#{1,6}\s+)", old_line)
                prefix = m.group(1) if m else ""
                new_line = prefix + repl
                lines[hit] = new_line + ("\n" if lines[hit].endswith("\n") else "")
                new_text = "".join(lines)
                changes.append({"para": "第 %d 行（标题）" % (hit + 1),
                                "before": old_line, "after": new_line,
                                "how": "整行替换标题"})
                msg = "整行替换标题"
                done = True
        if not done:
            if not find:
                return changes, False, "没有要处理的目标词（find 为空）", None
            n = text.count(find)
            if n == 0:
                return changes, False, "文件里找不到「%s」" % find, None
            new_text = text.replace(find, repl)
            changes.append({"para": "全文", "before": "「%s」×%d" % (find, n),
                            "after": "「%s」×%d" % (repl, n),
                            "how": "文本替换"})
            msg = "文本替换"
    elif op == "delete":
        if not find:
            return changes, False, "没有要删除的目标词", None
        new_text = text
        n = 0
        for _ in range(200):
            i = new_text.find(find)
            if i == -1:
                break
            after = new_text[i + len(find): i + len(find) + 1]
            if after in SEP_CHARS:
                new_text = new_text[:i] + new_text[i + len(find) + 1:]
            else:
                new_text = new_text[:i] + new_text[i + len(find):]
            n += 1
        if n == 0:
            return changes, False, "文件里找不到「%s」" % find, None
        changes.append({"para": "全文", "before": "删「%s」×%d" % (find, n),
                        "after": "已删", "how": "文本删除"})
        msg = "文本删除"
    elif op == "insert":
        if not add:
            return changes, False, "要加的内容为空", None
        lines = text.splitlines(keepends=True)
        hit = None
        for i, ln in enumerate(lines):
            if hint and hint in ln:
                hit = i
                break
        if hit is None:
            return changes, False, "没找到含「%s」的行" % hint, None
        ln = lines[hit].rstrip("\n")
        lines[hit] = ln + "、" + add + "\n"
        new_text = "".join(lines)
        changes.append({"para": "第 %d 行" % (hit + 1), "before": ln,
                        "after": lines[hit].rstrip("\n"), "how": "行尾追加"})
        msg = "行尾追加"
    else:
        return changes, False, "不认识的编辑操作：%s" % op, None

    return changes, True, msg, new_text


def edit_xlsx_cells(src: Path, out_path: Path, intent: dict):
    """对 xlsx 执行编辑：单元格文字替换 / 删除。返回 (改动明细, 是否成功, 说明)。

    设计要点：
    - **保格式**：只改 cell.value，样式（字体/边框/底色/列宽）天然不动
      （沿用写回器 fill.py 的经验：openpyxl 只改值 = 样式最省心路径）。
    - **只处理文本格**：数值格不动（防止把"120"误伤成"1200"的子串替换）。
    - **合并单元格**：附属格（MergedCell）不可写，跳过；锚点格正常处理。
    - 函数内直接保存到 out_path（openpyxl 的 load→改→save 一体，与 docx 分支
      "调用方保存"的差异在此说明）。
    """
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell

    op = intent.get("op")
    params = intent.get("params") or {}
    find = (params.get("find") or "").strip()
    repl = params.get("replace_with") or ""

    changes = []
    if op not in ("replace", "delete"):
        return changes, False, "xlsx 目前支持「替换 / 删除」两类编辑（收到：%s）" % op
    if not find:
        return changes, False, "没有要处理的目标词（find 为空）"

    wb = load_workbook(str(src))
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                v = cell.value
                if not isinstance(v, str) or find not in v:
                    continue
                old = v
                if op == "replace":
                    new = v.replace(find, repl)
                else:  # delete：删词 + 收拢残留分隔符
                    new = v.replace(find, "")
                    new = re.sub(r"[、，；,;]{2,}", "、", new).strip()
                    new = new.strip("、，；,; ")
                cell.value = new if new else None
                n += 1
                changes.append({
                    "para": "工作表「%s」· %s" % (ws.title, cell.coordinate),
                    "before": old,
                    "after": new,
                    "how": "单元格 %s（%s）" % (cell.coordinate, ws.title),
                })

    if n == 0:
        return changes, False, "整个工作簿里找不到「%s」" % find

    wb.save(str(out_path))
    verb = "替换" if op == "replace" else "删除"
    return changes, True, "命中 %d 个单元格（%s「%s」）" % (n, verb, find)


# ==================== 五、主流程 ====================

_QUOTED_RE = re.compile(r"['\"“”‘’「『]([^'\"“”‘’」』\n]{1,24})['\"“”‘’」』]")
_VERB_RE = re.compile(r"改成|改为|换成为?|变为|替换[为成]|变更为")


def _fill_edit_target(intent: dict, text: str):
    """兜底（模型飘移之一）：模型偶尔把引号里的目标词解析丢了——
    如"把周报里'本周完成'改成'本周进展'"被解析成 find 空、执行时报错。
    从指令原文的引号里把词捡回来：按"改成/换成/…"切分，动词前的引号词
    是"要找的词"（find），动词后的是"替换成的词"（replace_with）。

    只填空位、不覆盖模型已给的值；"把标题改成'X'"天然安全
    （动词前没有引号词 → find 保持空 → 标题分支照常接手，引号词补进 repl）。
    """
    if intent.get("action") != "编辑":
        return
    op = intent.get("op")
    params = intent.get("params")
    if not isinstance(params, dict):
        params = {}
        intent["params"] = params

    if op == "replace":
        m = _VERB_RE.search(text or "")
        head = text[:m.start()] if m else (text or "")
        tail = text[m.end():] if m else ""
        if not (params.get("find") or "").strip():
            w = _QUOTED_RE.findall(head)
            if w:
                params["find"] = w[0].strip()
        if not (params.get("replace_with") or "").strip():
            w = _QUOTED_RE.findall(tail)
            if w:
                params["replace_with"] = w[0].strip()
    elif op == "delete":
        if not (params.get("find") or "").strip():
            w = _QUOTED_RE.findall(text or "")
            if w:
                params["find"] = w[0].strip()
    elif op == "insert":
        if not (params.get("text") or "").strip():
            w = _QUOTED_RE.findall(text or "")
            if w:
                params["text"] = w[0].strip()


# 多操作（multi）用到的对照表：op → 所属类别 / 中文名
_OP_TO_ACTION = {
    "replace": "编辑", "insert": "编辑", "delete": "编辑",
    "set_align": "排版", "set_spacing": "排版", "set_indent": "排版",
    "set_bold": "格式调整", "set_italic": "格式调整", "set_underline": "格式调整",
    "set_strike": "格式调整", "set_highlight": "格式调整",
    "set_color": "格式调整", "set_font": "格式调整", "set_size": "格式调整",
}

_OP_CN = {
    "replace": "替换文字", "insert": "插入文字", "delete": "删除文字",
    "set_align": "对齐", "set_spacing": "行距", "set_indent": "缩进",
    "set_bold": "加粗", "set_italic": "斜体", "set_underline": "下划线",
    "set_strike": "删除线", "set_highlight": "高亮", "set_color": "颜色",
    "set_font": "字体", "set_size": "字号",
}


def _do_transform(action: str, intent: dict, working: Path, out_path: Path, text: str):
    """执行 编辑 / 排版 / 格式调整 的落地（不含提取）。
    返回 (ok, msg, changes)；调用方负责报告、清理与 meta 回填。"""
    suffix = working.suffix.lower()
    if action == "编辑":
        if suffix == ".docx":
            d = docx.Document(str(working))
            changes, ok2, msg = edit_docx_paragraphs(d, intent)
            if ok2:
                d.save(str(out_path))
            return ok2, msg, changes
        if suffix in (".md", ".txt"):
            changes, ok2, msg, new_text = edit_plain_text(working, intent)
            if ok2:
                out_path.write_text(new_text, encoding="utf-8")
            return ok2, msg, changes
        if suffix == ".xlsx":
            changes, ok2, msg = edit_xlsx_cells(working, out_path, intent)
            return ok2, msg, changes
        return False, "暂不支持 %s 类型的编辑（目前支持 docx / xlsx / md / txt）" % suffix, []
    # 排版 / 格式调整 → 交给 format_ops 执行
    changes, ok2, msg = apply_format(working, out_path, intent, text)
    return ok2, msg, changes


def _run_multi(text: str, intent: dict, src: Path, working: Path,
               report: list, meta: dict):
    """多操作（一句话里多个操作）：按顺序逐个执行，后一步基于前一步的产物继续；
    任何一步失败 → 删掉本次的中间产物，保证「要么都做了、要么都没做」（原件不动）。"""
    raw = intent.get("tasks")
    tasks = [t for t in raw if isinstance(t, dict) and t.get("op")] \
        if isinstance(raw, list) else []
    if not tasks:
        report.append("")
        report.append("【失败】解析成了「多操作」，但没给出具体操作清单（tasks 为空）——"
                      "请换个说法重发一次。")
        return False, "\n".join(report)
    for t in tasks:
        if t.get("op") == "extract":
            report.append("")
            report.append("【失败】多操作里暂不支持混合「内容提取」——"
                          "提取会产出清单文件，请单独发一条指令。")
            return False, "\n".join(report)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    first_action = _OP_TO_ACTION.get(tasks[0].get("op"), "格式调整")
    out_word = "_编辑结果" if first_action == "编辑" else "_格式结果"
    out_path = OUT_DIR / (src.stem + out_word + src.suffix)

    report.append("")
    report.append("【多操作】共 %d 个操作，按顺序执行（后一步基于前一步的结果）：" % len(tasks))

    cur = working
    wrote = False
    for i, t in enumerate(tasks, 1):
        op = t.get("op")
        sub_action = _OP_TO_ACTION.get(op)
        cn = _OP_CN.get(op, op)
        if sub_action is None:
            report.append("")
            report.append("【失败】第 %d 个操作不认识：%s" % (i, op))
            if wrote and out_path.exists():
                try:
                    out_path.unlink()
                except OSError:
                    pass
            return False, "\n".join(report)
        sub = {
            "action": sub_action, "op": op, "doc": src.name,
            "doc_hint": intent.get("doc_hint") or "",
            "locate": (t.get("locate") if isinstance(t.get("locate"), dict)
                       else {"by": "whole", "hint": ""}),
            "params": t.get("params") if isinstance(t.get("params"), dict) else {},
        }
        report.append("")
        report.append("── 第 %d / %d 步：%s（%s） ──" % (i, len(tasks), cn, op))
        try:
            ok2, msg, changes = _do_transform(sub_action, sub, cur, out_path, text)
        except PermissionError:
            ok2, msg, changes = False, (
                "输出文件正被别的程序占用（可能你正开着它）——"
                "请关掉「%s」再跑一次。" % out_path.name), []
        if not ok2:
            if wrote and out_path.exists():
                try:
                    out_path.unlink()
                    report.append("")
                    report.append("【回滚】第 %d 步失败，已删掉中间产物「%s」——"
                                  "为保证「要么都做、要么都不做」，本次不留半成品"
                                  "（原件未动）。" % (i, out_path.name))
                except OSError:
                    pass
            report.append("【失败】第 %d 步（%s）：%s" % (i, cn, msg))
            return False, "\n".join(report)
        wrote = True
        cur = out_path                      # 下一步从这个产物继续
        report.append("【定位】%s" % msg)
        report.append("【执行】共 %d 处改动：" % len(changes))
        for c in changes:
            report.append("  · %s" % c["how"])
            report.append("    改前：%s" % (c["before"][:90] if isinstance(c["before"], str) else c["before"]))
            report.append("    改后：%s" % (c["after"][:90] if isinstance(c["after"], str) else c["after"]))

    report.append("")
    report.append("【输出】已生成：%s（%d 个操作的全部修改都在这一份里）" % (out_path, len(tasks)))
    report.append("【原件】未改动（打开的是副本，保存到上面这个新文件）")
    report.append("=" * 62)
    if meta is not None:
        meta["ok"] = True
        meta["out_path"] = str(out_path)
    return True, "\n".join(report)


def _auto_store(files):
    """把提取产物 json 落进数据库（模块2 第4步「结果入库」）。

    失败只返回一行提示，绝不拦主流程——数据库不可用时提取照常交付。
    """
    try:
        jp = next((str(f) for f in files if str(f).endswith(".json")), None)
        if not jp:
            return ""
        from store import save_any            # 延迟导入：不拖累其他路径启动
        r = save_any(jp)
        kind_cn = {"entities": "实体", "profile": "画像", "extract": "字段"}.get(r["kind"], r["kind"])
        extra = (" · 移除 %d 条" % r["removed"]) if r.get("removed") else ""
        return "【入库】已存入数据库（%s：新增 %d 条 · 更新 %d 条%s）" % (
            kind_cn, r["new"], r["updated"], extra)
    except Exception as e:
        msg = str(e).replace("\n", " ")
        if len(msg) > 80:
            msg = msg[:80] + "…"
        return "【入库】跳过（数据库不可用：%s）——不影响本次提取结果" % msg


def run_edit(text: str, docs_dir: Path, context: dict = None, meta: dict = None,
             workcopy: dict = None, image: str = None):
    """执行一条编辑指令的完整流程，返回 (是否成功, 报告文本)。

    context（可选）＝ 会话上下文（多轮指代消解），透传给解析器：
      {"current_doc": "上一句操作的文档名"}
    meta（可选）＝ 调用方传一个空 dict 进来，函数把本轮关键信息填进去：
      intent（解析结果）｜ doc（实际作用的文档名）｜ ok（成功与否）｜ out_path（输出文件）
      给会话层（chat.py）与数据库记录用；不传则行为与之前完全一致。
    workcopy（可选）＝ 会话的"工作副本"（多轮累积修改）：
      {"doc": "这份副本属于哪份文档", "path": 副本路径}
      仅当"本轮找到的文档 == 副本归属文档"时才生效——基于副本继续改，
      多轮指令的修改累积到同一份最新文件；换文档自动不生效（回到原件）。
    image（可选）＝ 待插入的图片引用（上传后的存储名 / 文件名 / 绝对路径）。
      只有 op=insert_image 时用；不传则自动取最近上传的一张。
    """
    if meta is not None:
        meta.setdefault("ok", False)

    report = []
    report.append("=" * 62)
    report.append("文档操作执行器 · 一句话 → 编辑 / 排版 / 格式 / 提取")
    report.append("=" * 62)
    report.append("指令：%s" % text)
    report.append("")

    # ---- ① 解析 ----
    doc_names = scan_docs(docs_dir)
    intent, cost, ok = parse_one(text, doc_names, context=context)
    if not ok:
        report.append("【解析】失败（网络或格式问题），请重跑一次。")
        return False, "\n".join(report)
    if meta is not None:
        meta["intent"] = intent
    report.append("【解析】（%.1fs）%s/%s ｜ 文档=%s" %
                  (cost, intent.get("action"), intent.get("op"), intent.get("doc")))
    report.append("  %s" % json.dumps(intent, ensure_ascii=False))

    action = intent.get("action")

    # 图片插入统一归到「编辑」类（模型偶尔会把 action 判成别的，op 才是准的）
    if intent.get("op") == "insert_image" and action != "不明确":
        action = "编辑"

    # 兜底：模型把引号里的目标词解析丢了 → 从指令原文捡回来（只补空位）
    _fill_edit_target(intent, text)

    # 多轮目标兜底（第 7 步新增）：本句用指代词（"把它的加粗去掉"）但没提具体目标
    # → 用上一轮的目标（context.last_hint）补 locate.hint。
    # 只在"hint 为空 + 句中有指代词"时补；说清了目标的句子不受影响。
    if context and context.get("last_hint"):
        loc = intent.get("locate")
        if not isinstance(loc, dict):
            loc = {}
            intent["locate"] = loc
        if not (loc.get("hint") or "").strip():
            ref = (intent.get("doc_hint") or "") + " " + (text or "")
            if any(k in ref for k in ("它", "这个", "这份", "该文档", "此文档",
                                      "刚才", "上述", "同一个")):
                loc["hint"] = context["last_hint"]

    global LAST_RUN_LABEL
    LAST_RUN_LABEL = {"编辑": "编辑", "排版": "格式", "格式调整": "格式",
                      "内容提取": "提取", "多操作": "格式"}.get(action, "执行")

    if action == "不明确":
        report.append("")
        report.append("【反问】指令信息不足，系统想先问清楚：")
        report.append("  " + (intent.get("clarify") or "（没给出反问内容）"))
        return False, "\n".join(report)
    if action not in ("编辑", "排版", "格式调整", "内容提取", "多操作"):
        report.append("")
        report.append("【提示】目前支持「编辑」「排版」「格式调整」「内容提取」四类操作"
                      "（一句话里多个操作也可以）；"
                      "「%s」不在其中。" % action)
        return False, "\n".join(report)

    # ---- 找目标文档 ----
    src = find_doc_file(intent.get("doc"), docs_dir)
    if src is None:
        src = find_doc_file(intent.get("doc_hint"), docs_dir)

    # 多轮兜底（第 5 步新增）：本句用的是指代词（"它/这份/没提具体文档"）
    # 且上一句有文档 → 就指向它（模型解析 + 规则兜底双保险）。
    # 注意：只在"没用具体称呼"时兜底；用户明确说"把周报的…"而库里没有周报时，
    # 仍按失败处理（不能悄悄改成别的文档）。
    if src is None and context and context.get("current_doc"):
        hint = (intent.get("doc_hint") or "").strip()
        is_pronoun = (hint == "" or any(k in hint for k in
                      ("它", "这个", "这份", "该文档", "此文档", "刚才", "上述", "同一个")))
        if is_pronoun:
            cand = find_doc_file(context["current_doc"], docs_dir)
            if cand is not None:
                src = cand
                report.append("（多轮指代：本句未提文档名，按上一句的《%s》处理）" % cand.name)

    if src is None:
        report.append("")
        if intent.get("op") == "insert_image":
            report.append("【失败】还不知道要插到哪份文档——请说清文档名，"
                          "例如：把这张图片插到《设备采购合同》的标题下面")
        else:
            report.append("【失败】文档库里找不到目标文档：%s" %
                          (intent.get("doc") or intent.get("doc_hint")))
        return False, "\n".join(report)
    if meta is not None:
        meta["doc"] = src.name
    report.append("")
    report.append("【文档】%s" % src)

    # 多轮累积（第 5 步新增）：会话传了工作副本、且副本归属正是本轮文档
    # → 就基于副本继续修改（多轮修改逐轮累积到同一份最新文件；原件始终不动）。
    # 归属校验放在这里做：切到别的文档时副本自动不生效，回到原件从零开始。
    working = src
    if workcopy and workcopy.get("path"):
        wc_path = Path(workcopy["path"])
        if workcopy.get("doc") == src.name and wc_path.exists():
            working = wc_path
            report.append("（多轮累积：基于上一轮产物《%s》继续修改）" % wc_path.name)

    # ---- 图片插入（op=insert_image）：图片 → 文档，自动定位置 ----
    if intent.get("op") == "insert_image":
        from image_ops import insert_image, resolve_image, latest_image   # 延迟导入
        img = resolve_image(image) if image else None
        if img is None and not image:
            # 接口没带图片：退而用最近上传的一张（CLI 也能直接用"最近那张图"）
            img = latest_image()
            if img is not None:
                report.append("（指令没指定图片，使用最近上传的《%s》）" % img.name)
        if img is None:
            report.append("")
            report.append("【失败】没有可插入的图片——请先在操作台点「插入图片」上传，"
                          "或在指令里给出图片文件名。")
            return False, "\n".join(report)

        report.append("")
        report.append("【图片】%s" % img)
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        out_path = OUT_DIR / (src.stem + "_编辑结果" + src.suffix)
        try:
            changes, ok2, msg = insert_image(working, out_path, img, intent)
        except PermissionError:
            report.append("【失败】输出文件正被别的程序占用（可能你正开着它）——"
                          "请关掉「%s」再跑一次。" % out_path.name)
            return False, "\n".join(report)

        if not ok2:
            report.append("")
            report.append("【失败】%s" % msg)
            return False, "\n".join(report)

        report.append("")
        report.append("【定位】%s" % msg)
        report.append("")
        report.append("【执行】共 %d 处改动：" % len(changes))
        for c in changes:
            report.append("  · %s" % c["how"])
            report.append("    改前：%s" % (c["before"][:90] if isinstance(c["before"], str) else c["before"]))
            report.append("    改后：%s" % (c["after"][:90] if isinstance(c["after"], str) else c["after"]))
        report.append("")
        report.append("【输出】已生成：%s" % out_path)
        report.append("【原件】未改动（打开的是副本，保存到上面这个新文件）")
        report.append("=" * 62)
        if meta is not None:
            meta["ok"] = True
            meta["out_path"] = str(out_path)
        return True, "\n".join(report)

    # ---- 多操作（一句话里多个操作）：链式执行（第 7 步新增）----
    if action == "多操作" or intent.get("op") == "multi":
        return _run_multi(text, intent, src, working, report, meta)

    # ---- 内容提取：只读操作（extract=用户指定字段 / profile=画像 / entities=实体数据）----
    if action == "内容提取":
        op3 = intent.get("op")
        if op3 == "profile":
            from summary_ops import run_profile     # 延迟导入：不拖累编辑/格式的启动
            ok3, sub_report, _files = run_profile(text, intent, src, OUT_DIR)
        elif op3 == "entities":
            from entity_ops import run_entities     # 延迟导入：不拖累编辑/格式的启动
            ok3, sub_report, _files = run_entities(text, intent, src, OUT_DIR)
        else:
            from extract_ops import run_extract     # 延迟导入：不拖累编辑/格式的启动
            ok3, sub_report, _files = run_extract(text, intent, src, OUT_DIR)
        report.append("")
        report.append(sub_report)
        # ---- 自动入库（模块2 第4步）：产物 json 落入数据库，失败不拦主流程 ----
        if ok3 and _files:
            _line = _auto_store(_files)
            if _line:
                report.append("")
                report.append(_line)
        if meta is not None:
            meta["ok"] = ok3
        return ok3, "\n".join(report)

    # ---- ② 定位 + ③ 执行 ----
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out_word = "_编辑结果" if action == "编辑" else "_格式结果"
    out_path = OUT_DIR / (src.stem + out_word + src.suffix)

    try:
        ok2, msg, changes = _do_transform(action, intent, working, out_path, text)
    except PermissionError:
        report.append("【失败】输出文件正被别的程序占用（可能你正开着它）——"
                      "请关掉「%s」再跑一次。" % out_path.name)
        return False, "\n".join(report)

    if not ok2:
        report.append("")
        report.append("【失败】%s" % msg)
        return False, "\n".join(report)

    report.append("")
    report.append("【定位】%s" % msg)
    report.append("")
    report.append("【执行】共 %d 处改动：" % len(changes))
    for c in changes:
        report.append("  · %s" % c["how"])
        report.append("    改前：%s" % (c["before"][:90] if isinstance(c["before"], str) else c["before"]))
        report.append("    改后：%s" % (c["after"][:90] if isinstance(c["after"], str) else c["after"]))

    report.append("")
    report.append("【输出】已生成：%s" % out_path)
    report.append("【原件】未改动（打开的是副本，保存到上面这个新文件）")
    report.append("=" * 62)
    if meta is not None:
        meta["ok"] = True
        meta["out_path"] = str(out_path)
    return True, "\n".join(report)


def main():
    ap = argparse.ArgumentParser(description="模块（1）· 编辑执行器（一句话 → 文档真的变了）")
    ap.add_argument("text", help="一条编辑指令")
    ap.add_argument("--docs", default=str(DEFAULT_DOCS_DIR), help="文档库目录")
    args = ap.parse_args()

    ok, report = run_edit(args.text, Path(args.docs))
    print(report)

    # 把报告存盘（覆盖式，保留最近一次；编辑 → 编辑报告.txt，格式 → 格式报告.txt）
    try:
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        (OUT_DIR / (LAST_RUN_LABEL + "报告.txt")).write_text(report, encoding="utf-8")
    except Exception:
        pass
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
