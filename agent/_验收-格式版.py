# -*- coding: utf-8 -*-
"""
_验收-格式版.py —— 模块（1）第 3 步验收：8 条格式/排版指令逐条真跑 + 归档 + 验证

流程：①记录 8 个源文件指纹 → ②逐条调 edit.py 真跑 → ③每条产出归档改名
      → ④读回输出文件逐条验证格式属性 → ⑤核对源文件指纹（原件零改动）

跑法（项目根目录）：
  python agent/_验收-格式版.py
"""
import hashlib
import os
import subprocess
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parents[1]
PY = sys.executable
AGENT = BASE / "agent"
OUT = AGENT / "输出"
DOCS = BASE / "测试数据" / "源文档"


def md5(p: Path) -> str:
    return hashlib.md5(p.read_bytes()).hexdigest()


# (编号, 指令, 输出文件名, 归档文件名)
CASES = [
    (5, "把设备采购合同的标题居中",
     "设备采购合同_格式结果.docx", "设备采购合同_格式结果_标题居中.docx"),
    (6, "把年度数字化转型工作总结的正文设置成 1.5 倍行距",
     "年度数字化转型工作总结_格式结果.docx", "年度数字化转型工作总结_格式结果_1.5倍行距.docx"),
    (7, "把研发项目立项管理办法的正文每段首行缩进两个字符",
     "研发项目立项管理办法_格式结果.docx", "研发项目立项管理办法_格式结果_首行缩进.docx"),
    (8, "把项目推进会会议纪要的会议时间那一行右对齐",
     "项目推进会会议纪要_格式结果.docx", "项目推进会会议纪要_格式结果_右对齐.docx"),
    (9, "把设备采购合同里所有的金额加粗",
     "设备采购合同_格式结果.docx", "设备采购合同_格式结果_金额加粗.docx"),
    (10, "把员工信息表的表头字号调成 14 号",
     "员工信息表_格式结果.xlsx", "员工信息表_格式结果_表头字号14.xlsx"),
    (11, "把项目推进会会议纪要的标题改成黑体",
     "项目推进会会议纪要_格式结果.docx", "项目推进会会议纪要_格式结果_标题黑体.docx"),
    (12, "把项目周报里所有的'问题'标红",
     "项目周报-第12周_格式结果.md", "项目周报-第12周_格式结果_问题标红.md"),
]


def main():
    print("=" * 66)
    print("第 3 步验收 · 8 条格式/排版指令逐条真跑")
    print("=" * 66)

    # ---- ① 源文件指纹 ----
    src_map = {
        5: DOCS / "docx" / "设备采购合同.docx",
        6: DOCS / "docx" / "年度数字化转型工作总结.docx",
        7: DOCS / "docx" / "研发项目立项管理办法.docx",
        8: DOCS / "docx" / "项目推进会会议纪要.docx",
        9: DOCS / "docx" / "设备采购合同.docx",
        10: DOCS / "xlsx" / "员工信息表.xlsx",
        11: DOCS / "docx" / "项目推进会会议纪要.docx",
        12: DOCS / "md" / "项目周报-第12周.md",
    }
    hashes_before = {k: md5(v) for k, v in src_map.items()}

    # ---- ② 逐条跑 + ③ 归档 ----
    results = {}
    for cid, text, outname, archname in CASES:
        print("\n---- 第 %d 条：%s" % (cid, text))
        r = subprocess.run([PY, str(AGENT / "edit.py"), text],
                           capture_output=True, text=True, encoding="utf-8",
                           errors="replace", timeout=240, cwd=str(BASE))
        ok = (r.returncode == 0)
        (OUT / ("第%d条-报告.txt" % cid)).write_text(r.stdout + "\n" + r.stderr, encoding="utf-8")
        o = OUT / outname
        archived = None
        if o.exists():
            archived = OUT / archname
            os.replace(str(o), str(archived))
        results[cid] = {"ok": ok, "out": archived, "stdout": r.stdout}
        tag = "OK " if ok else "FAIL"
        print("  [%s] 退出码=%d ｜ 产物=%s" % (tag, r.returncode,
              archname if archived else "（无产物）"))

    # ---- ④ 逐个验证 ----
    print("\n" + "=" * 66)
    print("逐条验证（读回输出文件，检查格式属性）")
    print("=" * 66)
    checks = []

    def check(cid, name, cond, detail):
        checks.append((cid, name, bool(cond), detail))

    # 5 · 标题居中
    o = results[5]["out"]
    if o:
        d = docx.Document(str(o))
        a0 = d.paragraphs[0].paragraph_format.alignment
        a6 = d.paragraphs[6].paragraph_format.alignment  # Heading 不该被动
        check(5, "标题段落对齐=CENTER", a0 == WD_ALIGN_PARAGRAPH.CENTER, "align=%s" % a0)
        check(5, "非目标段落对齐未变", a6 is None, "段6 align=%s" % a6)

    # 6 · 1.5 倍行距
    o = results[6]["out"]
    if o:
        d = docx.Document(str(o))
        ls = [d.paragraphs[i].paragraph_format.line_spacing for i in (5, 8, 13)]
        ls_head = d.paragraphs[4].paragraph_format.line_spacing
        check(6, "正文行距=1.5（抽3段）", all(x == 1.5 for x in ls), "spacing=%s" % ls)
        check(6, "标题行距未变", ls_head is None, "段4 spacing=%s" % ls_head)

    # 7 · 首行缩进
    o = results[7]["out"]
    if o:
        d = docx.Document(str(o))
        def flc(i):
            pPr = d.paragraphs[i]._p.find(qn("w:pPr"))
            ind = pPr.find(qn("w:ind")) if pPr is not None else None
            return ind.get(qn("w:firstLineChars")) if ind is not None else None
        vals = [flc(i) for i in (5, 9, 20)]
        check(7, "正文首行缩进=200（抽3段）", all(v == "200" for v in vals), "chars=%s" % vals)

    # 8 · 右对齐
    o = results[8]["out"]
    if o:
        d = docx.Document(str(o))
        tgt = next(p for p in d.paragraphs if "会议时间" in p.text)
        nxt = next(p for p in d.paragraphs if "会议地点" in p.text)
        check(8, "「会议时间」段=RIGHT", tgt.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.RIGHT,
              "align=%s" % tgt.paragraph_format.alignment)
        check(8, "「会议地点」段未变", nxt.paragraph_format.alignment is None,
              "align=%s" % nxt.paragraph_format.alignment)

    # 9 · 金额加粗
    o = results[9]["out"]
    if o:
        d = docx.Document(str(o))
        p = d.paragraphs[9]
        bt = [r.text for r in p.runs if r.bold]
        full = "".join(r.text for r in p.runs)
        orig = "本合同总金额为人民币壹佰捌拾万元整（￥1,800,000.00）。"
        check(9, "加粗文本=2处金额", bt == ["人民币壹佰捌拾万元整", "￥1,800,000.00"], "bold=%s" % bt)
        check(9, "段落文本与原文一致", full == orig, "拼接=%r" % full[:40])
        nb = sum(1 for r in d.paragraphs[10].runs if r.bold)
        check(9, "邻段零误伤", nb == 0, "段10 加粗run=%d" % nb)

    # 10 · 表头字号 14
    o = results[10]["out"]
    if o:
        wb = load_workbook(str(o))
        ws = wb.active
        sizes_h = [ws.cell(row=2, column=c).font.size for c in range(1, 13)]
        sizes_d = [ws.cell(row=3, column=c).font.size for c in range(1, 13)]
        check(10, "表头全 14 号", all(s == 14 for s in sizes_h), "sizes=%s" % set(sizes_h))
        check(10, "数据行仍 11 号（零误伤）", all(s == 11 for s in sizes_d), "sizes=%s" % set(sizes_d))
        wb.close()

    # 11 · 标题黑体
    o = results[11]["out"]
    if o:
        d = docx.Document(str(o))
        r0 = d.paragraphs[0].runs[0]
        rPr = r0._element.rPr
        rFonts = rPr.rFonts if rPr is not None else None
        ea = rFonts.get(qn("w:eastAsia")) if rFonts is not None else None
        check(11, "标题字体=黑体（含 eastAsia）", r0.font.name == "黑体" and ea == "黑体",
              "name=%s eastAsia=%s" % (r0.font.name, ea))
        r1name = d.paragraphs[1].runs[0].font.name if d.paragraphs[1].runs else None
        check(11, "正文未变黑体", r1name != "黑体", "段1 font=%s" % r1name)

    # 12 · md 标红
    o = results[12]["out"]
    if o:
        t = o.read_text(encoding="utf-8")
        n = t.count('<span style="color:#FF0000">问题</span>')
        t_clean = t.replace('<span style="color:#FF0000">问题</span>', "问题")
        orig = (DOCS / "md" / "项目周报-第12周.md").read_text(encoding="utf-8")
        check(12, "5 处「问题」全部包红", n == 5, "span 数=%d" % n)
        check(12, "剥离标签后与原文一致", t_clean == orig, "len=%d vs %d" % (len(t_clean), len(orig)))

    # ---- ⑤ 源文件指纹核对 ----
    print()
    src_ok = True
    for cid, v in src_map.items():
        if md5(v) != hashes_before[cid]:
            src_ok = False
            print("  [!!] 源文件被改动：%s" % v)
    print("[源文件] 8 个源文件指纹 %s" % ("全部未改动 ✓" if src_ok else "有改动 ✗"))

    # ---- 汇总 ----
    print()
    print("=" * 66)
    n_pass = sum(1 for _, _, ok, _ in checks if ok)
    for cid, name, ok, detail in checks:
        print("  [%s] 第%d条 %s ｜ %s" % ("PASS" if ok else "FAIL", cid, name, detail))
    print("-" * 66)
    print("小计：%d/%d 项验证通过" % (n_pass, len(checks)))
    run_ok = sum(1 for c in results.values() if c["ok"])
    print("执行：%d/8 条指令运行成功" % run_ok)
    print("=" * 66)
    return 0 if (n_pass == len(checks) and run_ok == 8 and src_ok) else 1


if __name__ == "__main__":
    sys.exit(main())
