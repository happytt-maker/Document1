# -*- coding: utf-8 -*-
"""
_验收-会话版.py —— 模块（1）· 第 5 步（交互循环 + MySQL）自动验收

一条命令复跑全套：
  1. 前置：数据库就绪（建库建表）
  2. 跑 5 轮连续会话（新会话号），每轮核对"作用文档"是否解析正确：
       ①带全信息锁定纪要 → ②③用「它」（应仍指纪要）
       → ④切换到合同 → ⑤「它」应跟着变成合同
  3. 累积效果读回验证：
       纪要_编辑结果.docx = 标题居中 + 会议时间右对齐 + 李博文已删
       设备采购合同_格式结果.docx = 金额加粗 + 标题字号 16
  4. 数据库核对：本会话 5 条记录、字段齐全
  5. 源文档零改动：源文档库整树 md5 跑前跑后一致

跑法（在项目根目录）：
  python agent/_验收-会话版.py
"""

import hashlib
import sys
import time
from pathlib import Path

BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE / "agent"))
sys.stdout.reconfigure(errors="replace")

import docx                                                    # noqa: E402
import chat                                                    # noqa: E402
import db                                                      # noqa: E402

DOCS_DIR = BASE / "测试数据" / "源文档"
OUT_DIR = BASE / "agent" / "输出"
SESSION = "s-验收-" + time.strftime("%m%d-%H%M")

checks = []          # (名称, 通过?, 说明)


def check(name, ok, detail=""):
    checks.append((name, bool(ok), detail))
    print("  [%s] %s %s" % ("PASS" if ok else "FAIL", name, ("｜ " + detail) if detail else ""))


def md5_tree(root: Path):
    """整棵目录的 md5 指纹（逐文件 hash 后汇总）。"""
    h = hashlib.md5()
    for p in sorted(root.rglob("*")):
        if p.is_file():
            h.update(p.name.encode("utf-8"))
            h.update(str(p.stat().st_size).encode())
            h.update(hashlib.md5(p.read_bytes()).hexdigest().encode())
    return h.hexdigest()


def main():
    print("=" * 64)
    print("第 5 步验收：交互循环（多轮指代 + 累积）· MySQL 落库")
    print("会话号：%s" % SESSION)
    print("=" * 64)

    # ---- 0) 前置 ----
    print()
    print("【0】前置检查")
    db.ensure_db()
    check("数据库就绪（a23_agent.commands）", True)

    fp_before = md5_tree(DOCS_DIR)

    # ---- 1) 跑 5 轮会话 ----
    print()
    print("【1】跑 5 轮连续会话")
    sess = chat.Session(DOCS_DIR, session_id=SESSION)
    steps = [
        ("把项目推进会会议纪要的标题居中", "项目推进会会议纪要.docx"),
        ("把它的会议时间那一行改成右对齐", "项目推进会会议纪要.docx"),
        ("再把它里面的'李博文'删掉", "项目推进会会议纪要.docx"),
        ("把设备采购合同的金额加粗", "设备采购合同.docx"),
        ("把它的标题字号改成16", "设备采购合同.docx"),
    ]
    results = []
    for i, (text, expect_doc) in enumerate(steps, 1):
        ok, report, meta, elapsed = sess.run_turn(text)
        got = meta.get("doc")
        results.append((i, ok, got, elapsed))
        check("第 %d 轮执行成功（%.1fs）" % (i, elapsed), ok, "指令：%s" % text)
        check("第 %d 轮文档 = %s" % (i, expect_doc), got == expect_doc, "实际 = %s" % got)

    # ---- 2) 累积效果读回 ----
    print()
    print("【2】最终文件读回（多轮累积）")
    p = OUT_DIR / "项目推进会会议纪要_编辑结果.docx"
    if p.exists():
        d = docx.Document(str(p))
        p0 = d.paragraphs[0]
        p_time = next((x for x in d.paragraphs if "会议时间" in x.text), None)
        p_mem = next((x for x in d.paragraphs if "参会人员" in x.text), None)
        check("纪要：标题居中（第1轮）",
              p0.paragraph_format.alignment is not None and "CENTER" in str(p0.paragraph_format.alignment))
        check("纪要：会议时间右对齐（第2轮）",
              p_time is not None and p_time.paragraph_format.alignment is not None
              and "RIGHT" in str(p_time.paragraph_format.alignment))
        check("纪要：李博文已删（第3轮）", p_mem is not None and "李博文" not in p_mem.text,
              (p_mem.text[:40] if p_mem else ""))
    else:
        check("纪要_编辑结果.docx 存在", False)

    p2 = OUT_DIR / "设备采购合同_格式结果.docx"
    if p2.exists():
        d2 = docx.Document(str(p2))
        title = d2.paragraphs[0]
        money = d2.paragraphs[9]
        sz = title.runs[0].font.size if title.runs else None
        check("合同：标题字号 16（第5轮）", sz is not None and abs(sz.pt - 16) < 0.01,
              "实际 = %s" % (sz.pt if sz else None))
        check("合同：金额加粗保留（第4轮）", any(r.bold for r in money.runs),
              "加粗片段 = %s" % [r.text for r in money.runs if r.bold])
    else:
        check("设备采购合同_格式结果.docx 存在", False)

    # ---- 3) 数据库核对 ----
    print()
    print("【3】数据库核对")
    conn = db.get_conn()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT COUNT(*) FROM commands WHERE session_id=%s", (SESSION,))
            n = cur.fetchone()[0]
            cur.execute("SELECT doc, action, op, ok FROM commands WHERE session_id=%s ORDER BY turn", (SESSION,))
            rows = cur.fetchall()
    finally:
        conn.close()
    check("本会话写入 5 条记录", n == 5, "实际 = %d" % n)
    check("记录字段齐全（doc 非空 ×5）", all(r[0] for r in rows),
          " docs = %s" % [r[0] for r in rows])

    # ---- 4) 源文档零改动 ----
    print()
    print("【4】源文档零改动")
    fp_after = md5_tree(DOCS_DIR)
    check("源文档库整树指纹一致", fp_before == fp_after,
          "指纹 %s…%s" % (fp_before[:8], fp_after[:8]))

    # ---- 汇总 ----
    print()
    print("=" * 64)
    n_pass = sum(1 for _, ok, _ in checks if ok)
    print("验收结果：%d / %d 通过" % (n_pass, len(checks)))
    for name, ok, _ in checks:
        if not ok:
            print("  ✗ 未过：%s" % name)
    print("=" * 64)
    sys.exit(0 if n_pass == len(checks) else 1)


if __name__ == "__main__":
    main()
