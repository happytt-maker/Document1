# -*- coding: utf-8 -*-
"""
生成word模板.py —— 把 5 张 xlsx 模板转成 word（.docx）版样本
================================================================

背景：
    赛题评价方法原文：「再分别上传 5 个模版表格文件（word 或 excel）」
    —— 比赛当天给的模板可能是 word。项目此前只测过 xlsx 模板，
    需要 word 版样本用于「Word 模板兼容」的开发与验收。

本脚本做什么：
    逐张读 测试数据/模板表格/*.xlsx，把整张表（标题/说明栏/表头/待填区/
    结尾行）原样搬进一个 docx 表格，保存到本目录（模板表格-word/）。

    为什么"原样搬"：第一版样本先保持与 xlsx 完全同构，
    这样 word 读取器的输出可以直接和 xlsx 版结构对照（同构=可对比）。
    （标题在外面的变体形态，后续需要时再加。）

用法：
    python 生成word模板.py          # 在 VSCode 里对着本文件按 F5
"""

import sys
from pathlib import Path

import openpyxl
from docx import Document

HERE = Path(__file__).resolve().parent          # 测试数据/模板表格-word/
BASE = HERE.parent.parent                        # 研一项目/
SRC = BASE / "测试数据" / "模板表格"              # xlsx 模板目录


def grid_of(xlsx_path):
    """读一张 xlsx 模板的二维格子表（空串算 None；与 template_reader 同口径）。"""
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True)
    ws = wb.worksheets[0]
    grid = []
    for row in ws.iter_rows():
        line = []
        for cell in row:
            v = cell.value
            if v is None or (isinstance(v, str) and v.strip() == ""):
                line.append(None)
            else:
                line.append(str(v).strip() if isinstance(v, str) else v)
        grid.append(line)
    wb.close()
    return grid


def make_docx(grid, out_path):
    """把二维格子表原样搬进一个 docx 表格。"""
    doc = Document()

    n_rows = len(grid)
    n_cols = max(len(r) for r in grid)
    table = doc.add_table(rows=n_rows, cols=n_cols)
    try:
        table.style = "Table Grid"       # 带边框，更像真实的表格模板
    except Exception:
        pass                             # 样式名不存在也不影响内容

    for ri, line in enumerate(grid):
        for ci in range(n_cols):
            v = line[ci] if ci < len(line) else None
            if v is not None:
                table.cell(ri, ci).text = str(v)

    doc.save(str(out_path))


def main():
    xlsx_files = sorted(SRC.glob("*.xlsx"))
    if not xlsx_files:
        print("没找到 xlsx 模板：%s" % SRC)
        return 1

    made = []
    for xlsx in xlsx_files:
        grid = grid_of(xlsx)
        out_docx = HERE / (xlsx.stem + ".docx")
        make_docx(grid, out_docx)

        n_filled = sum(1 for line in grid for v in line if v is not None)
        n_rows = len(grid)
        n_cols = max(len(r) for r in grid)
        made.append((out_docx.name, n_rows, n_cols, n_filled))

    print("生成 word 版模板 %d 张：" % len(made))
    for name, r, c, f in made:
        print("  %s  （%d 行 × %d 列，有内容格 %d）" % (name, r, c, f))
    print("目录：%s" % HERE)
    return 0


if __name__ == "__main__":
    sys.exit(main())
