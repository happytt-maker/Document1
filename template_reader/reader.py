# -*- coding: utf-8 -*-
"""
template_reader/reader.py
-------------------------
模板表结构读取器。

目标：给它一张陌生的空白模板表，它输出一份「填表说明书」：
    - 表头在哪一行
    - 每一列叫什么
    - 哪些位置需要填（待填区域的范围）

═══════════════════════════════════════════════════════════════════
  和前一个项目（doc_reader）的关系
═══════════════════════════════════════════════════════════════════

doc_reader 解决的是「源文档里有什么」——把真实资料读成内容块。

这个项目的目标不同：读的是「空白的表格」，回答的是
「这张表想让我填什么」。所以不需要读公式、不需要读合并单元格
的重复内容，重点是**判断行和列的角色**。

所以这里没有复用 doc_reader 的读取器，而是重写了一个更简单的 ——
只读每个格子的「值」和「是不是空的」，够用了。

═══════════════════════════════════════════════════════════════════
  核心难点：表头行怎么找？
═══════════════════════════════════════════════════════════════════

模板表通常长这样：

    行1   设备采购汇总表                    ← 大标题（1格，常是合并）
    行2   填表单位： | 填表日期：            ← 说明栏（2格）
    行3   合同编号： | 合同金额：            ← 说明栏（2格）
    行4   序号|设备名称|规格型号|数量|单价|供应商  ← 真正的表头
    行5-12  （全空）                          ← 待填区
    行13  合计金额                            ← 结尾

【踩过的坑】一开始想用「第一个非空格子数 ≥ 2 的行」当表头。
结果行2、行3 也是 2 个格子，抢在表头前面，判断失败。

【正确做法】找「非空格子数最多」的那一行。
    实测：模板1 表头行 6 格、模板2 表头行 8 格 —— 都等于这一表的列数，
    明显比说明栏的 2 格多。

【辅助判据】表头行的文字短、不带「：」；说明栏的文字带「：」。
    两条一起用，判断更稳。

═══════════════════════════════════════════════════════════════════
"""

from pathlib import Path

import openpyxl
from docx import Document


# -----------------------------------------------------------------
# 第一步：把一个工作表读成「二维格子表」
# -----------------------------------------------------------------

def read_grid(sheet):
    """
    把工作表读成一个二维列表，方便按行列操作。

    返回 [[值或None, 值或None, ...], ...]
        外层是行，内层是列。坐标从 0 开始（grid[0] 是第 1 行）。

    空字符串会被当成 None（很多模板表的空格子是空字符串，不是真的 None）。
    """
    grid = []
    for row in sheet.iter_rows():
        line = []
        for cell in row:
            v = cell.value
            # 空字符串、纯空格都算"空"
            if v is None:
                line.append(None)
            elif isinstance(v, str) and v.strip() == "":
                line.append(None)
            else:
                line.append(str(v).strip() if isinstance(v, str) else v)
        grid.append(line)
    return grid


def count_filled(row):
    """数这一行有几个非空格子。"""
    return sum(1 for v in row if v is not None)


# -----------------------------------------------------------------
# 第二步：找表头行
# -----------------------------------------------------------------

def find_header_row(grid):
    """
    找表头行，返回行号（从 1 开始）。

    规则：非空格子数最多的那一行。
    如果有多行格子数一样多，取最靠上的那一行。

    找不到（整张表都是空的）就返回 None。
    """
    best_row = None
    best_count = 0

    for idx, row in enumerate(grid, start=1):
        n = count_filled(row)

        # 至少要有 2 个格子才可能是表头（1 格的是大标题）
        if n < 2:
            continue

        if n > best_count:
            best_count = n
            best_row = idx

    return best_row


def header_looks_right(grid, header_row):
    """
    额外检查：表头行的文字是不是"不像说明栏"。

    说明栏的特征：文字里带「：」
    表头行的特征：文字短、不带「：」

    返回 (是否像表头, 原因说明)
    这个函数不改变判断结果，只是给个提示，方便人工核对。
    """
    if header_row is None:
        return False, "没有找到候选表头行"

    row = grid[header_row - 1]
    texts = [str(v) for v in row if v is not None]

    colon_count = sum(1 for t in texts if "：" in t or ":" in t)
    if colon_count > 0:
        return False, "有 {0} 个格子的文字带冒号，像说明栏不像表头".format(colon_count)

    return True, "文字不带冒号，符合表头特征"


# -----------------------------------------------------------------
# 第三步：提取列标题
# -----------------------------------------------------------------

def read_columns(grid, header_row):
    """
    按表头行取出每一列的标题。

    返回 [{"col": 1, "title": "序号"}, {"col": 2, "title": "设备名称"}, ...]

    表头行是空格子的列也会记录，title 记为 None ——
    因为下游要按"列的位置"对齐，不能因为没标题就跳过（会错位）。
    """
    if header_row is None:
        return []

    row = grid[header_row - 1]
    columns = []
    for col_idx, v in enumerate(row, start=1):
        columns.append({
            "col": col_idx,
            "title": str(v) if v is not None else None,
        })
    return columns


# -----------------------------------------------------------------
# 第四步：找待填区域
# -----------------------------------------------------------------

def find_fill_region(grid, header_row):
    """
    找"要被填进去的那片空白区"。

    规则：表头行的下一行开始，往下找**第一段连续的全空行**。
         遇到一个有内容的行就停（那通常是个结尾行，比如「合计金额：」）。

    返回 {"start_row": 5, "end_row": 12, "row_count": 8}
    找不到就返回 None。
    """
    if header_row is None:
        return None

    start = None
    end = None

    for idx in range(header_row + 1, len(grid) + 1):
        row = grid[idx - 1]
        is_empty = count_filled(row) == 0

        if is_empty:
            if start is None:
                start = idx
            end = idx
        else:
            # 遇到有内容的行：如果已经开始记录空白区，就到此结束
            if start is not None:
                break
            # 如果还没开始（表头紧跟着有内容的行），继续往下找

    if start is None:
        return None

    return {
        "start_row": start,
        "end_row": end,
        "row_count": end - start + 1,
    }


# -----------------------------------------------------------------
# 第五步：读一张表（把上面几步串起来）
# -----------------------------------------------------------------

def read_template(xlsx_path):
    """
    读一张模板表，返回完整的「填表说明书」。

    返回结构：
        {
          "file": "模板1-设备采购汇总表.xlsx",
          "sheet_count": 1,
          "sheets": [
            {
              "name": "设备采购汇总",
              "max_row": 14,
              "max_column": 6,
              "header_row": 4,
              "header_ok": True,
              "header_note": "文字不带冒号，符合表头特征",
              "columns": [{"col":1,"title":"序号"}, ...],
              "fill_region": {"start_row":5, "end_row":12, "row_count":8},
            }
          ]
        }
    """
    xlsx_path = Path(xlsx_path)
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True)

    sheets = []
    for sheet in wb.worksheets:
        grid = read_grid(sheet)

        header_row = find_header_row(grid)
        header_ok, note = header_looks_right(grid, header_row)
        columns = read_columns(grid, header_row)
        fill_region = find_fill_region(grid, header_row)

        sheets.append({
            "name": sheet.title,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "header_row": header_row,
            "header_ok": header_ok,
            "header_note": note,
            "columns": columns,
            "fill_region": fill_region,
        })

    wb.close()

    return {
        "file": xlsx_path.name,
        "format": "xlsx",
        "sheet_count": len(sheets),
        "sheets": sheets,
    }


# -----------------------------------------------------------------
# docx（Word）模板支持（2026-10-14 新增）
# -----------------------------------------------------------------
# 赛题评价方法原文：「再分别上传 5 个模版表格文件（word 或 excel）」
# —— 模板可能是 Word。Word 模板 = docx 里的一张或多张表格。
#
# 【关键设计：共用 grid 逻辑】
#   表头行 / 列 / 待填区的判断（find_header_row / read_columns /
#   find_fill_region）本来就只依赖"二维格子表"，与文件格式无关。
#   所以这里只做一件事：把 docx 表格转成同样的 grid，后面全部复用。
#
# 【docx 表格的坑（沿用 doc_reader 的实测经验）】
#   1. row.cells 对合并单元格会重复返回同一格（包装对象不同、id 不可靠）
#      → 用 cell._tc 去重
#   2. 合并格的 text 会叠加（'吴晓敏\n吴晓敏'）→ 相同行只留一份

def read_docx_grid(table):
    """把一个 python-docx 表格转成二维格子表（与 read_grid 同口径：空格记 None）。

    注意：合并单元格按"去重后的格子顺序"排，行长度可能不一致
    （除表头行外的行不参与列判断，不影响后续逻辑）。
    """
    grid = []
    for row in table.rows:
        line = []
        seen = set()
        for cell in row.cells:
            key = cell._tc
            if key in seen:
                continue
            seen.add(key)

            v = (cell.text or "").strip()
            # 合并格叠加去重：所有行相同 → 留一份；不同 → 空格连接
            if "\n" in v:
                parts = [p.strip() for p in v.split("\n") if p.strip()]
                if parts and len(set(parts)) == 1:
                    v = parts[0]
                elif parts:
                    v = " ".join(parts)
            line.append(v if v else None)
        grid.append(line)
    return grid


def read_docx_template(docx_path):
    """读一张 docx（Word）模板，输出与 read_template 同构的「填表说明书」。

    每个表格当「一个工作表」处理，name 记为「表格1」「表格2」…
    （下游按顺序索引即可对应回 doc.tables）
    """
    docx_path = Path(docx_path)
    doc = Document(str(docx_path))

    sheets = []
    for ti, table in enumerate(doc.tables, start=1):
        grid = read_docx_grid(table)

        header_row = find_header_row(grid)
        header_ok, note = header_looks_right(grid, header_row)
        columns = read_columns(grid, header_row)
        fill_region = find_fill_region(grid, header_row)

        sheets.append({
            "name": "表格%d" % ti,
            "max_row": len(grid),
            "max_column": max((len(r) for r in grid), default=0),
            "header_row": header_row,
            "header_ok": header_ok,
            "header_note": note,
            "columns": columns,
            "fill_region": fill_region,
        })

    return {
        "file": docx_path.name,
        "format": "docx",
        "sheet_count": len(sheets),
        "sheets": sheets,
    }


# -----------------------------------------------------------------
# 方便直接运行：读一个文件夹里的所有模板表（xlsx + docx 都读）
# -----------------------------------------------------------------

def read_folder(folder):
    """读一个文件夹里所有模板（xlsx 和 docx），返回说明书列表。"""
    folder = Path(folder)
    results = []

    files = sorted(list(folder.glob("*.xlsx")) + list(folder.glob("*.docx")))
    for p in files:
        try:
            if p.suffix.lower() == ".docx":
                results.append(read_docx_template(p))
            else:
                results.append(read_template(p))
        except Exception as e:
            results.append({
                "file": p.name,
                "error": repr(e),
            })

    return results
